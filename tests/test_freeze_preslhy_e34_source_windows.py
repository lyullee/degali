from pathlib import Path
import json
import runpy

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
TOOL = runpy.run_path(ROOT / "tools/freeze_preslhy_e34_source_windows.py")


def _source_only_workbook(path: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "TE-WG"
    sheet.cell(7, 1, "X_Value")
    sheet.cell(7, 31, "Sync. Time [s]")
    sheet.cell(7, 33, "m(LH2) [g]")
    for row, time in enumerate(range(501), 8):
        if 40 <= time <= 60:
            mass = 100.0 * (time - 40)
        elif 60 < time <= 200:
            mass = 2000.0 - (time - 60) * (1980.0 / 140.0)
        elif 240 <= time <= 260:
            mass = 75.0 * (time - 240)
        elif 260 < time <= 450:
            mass = 1500.0 - (time - 260) * (1480.0 / 190.0)
        else:
            mass = 0.0
        sheet.cell(row, 1, time)
        sheet.cell(row, 31, time)
        sheet.cell(row, 33, mass)
    # Deliberately omit the dynamic H2/wind sheet.  The freeze must not need it.
    book.save(path)


def test_build_freeze_uses_only_source_mass_and_retains_all_windows(tmp_path):
    workbook = tmp_path / "20200417-Concrete03Wind-Final.xlsx"
    output = tmp_path / "freeze"
    _source_only_workbook(workbook)

    record = TOOL["build_source_window_freeze"](tmp_path, output)

    assert record["trial_count"] == 1
    assert record["total_window_count"] == 2
    assert record["selection_frozen_before_model_scoring"] is True
    assert record["all_qualifying_windows_retained"] is True
    assert record["selection_input_channels"] == [
        "Sync. Time [s]", "m(LH2) [g]",
    ]
    assert record["atmospheric_vapour_flux_qualified"] is False
    assert (output / "source-window-freeze-review.png").is_file()
    saved = json.loads((output / "source-window-freeze.json").read_text())
    assert saved["total_window_count"] == 2
    assert all(
        window["source_only_selection"]
        for window in saved["trials"][0]["windows"]
    )
