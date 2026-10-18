from pathlib import Path
import json
import math
import runpy

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
FREEZE_TOOL = runpy.run_path(
    ROOT / "tools/freeze_preslhy_e34_source_windows.py"
)
AUDIT_TOOL = runpy.run_path(
    ROOT / "tools/audit_preslhy_e34_dynamic_evidence.py"
)
REPLAY_TOOL = runpy.run_path(
    ROOT / "tools/build_preslhy_e34_dynamic_replay_package.py"
)


def _workbook(path: Path) -> None:
    book = openpyxl.Workbook()
    pool = book.active
    pool.title = "TE-WG"
    pool.cell(7, 1, "X_Value")
    pool.cell(7, 2, "TA000-00")
    pool.cell(7, 31, "Sync. Time [s]")
    pool.cell(7, 33, "m(LH2) [g]")
    for row, time in enumerate(range(301), 8):
        if 20 <= time <= 40:
            mass = 100.0 * (time - 20)
        elif 40 < time <= 220:
            mass = 2000.0 - (time - 40) * (1980.0 / 180.0)
        else:
            mass = 0.0
        pool.cell(row, 1, time)
        pool.cell(row, 2, 293.0)
        pool.cell(row, 31, time)
        pool.cell(row, 33, mass)

    dynamic = book.create_sheet("cH2-Wind")
    headers = (
        "Orig.Time [s]", "Sync. Time [s]", "CH35", "CH45", "CH55",
        "Young_T", "Young_U", "Young_V", "Young_W", "Trigger",
        "cH2 Time [s]",
    )
    for column, label in enumerate(headers, 1):
        dynamic.cell(6, column, label)
        dynamic.cell(5, column, 0.1)
    dynamic["C5"] = "defective"
    for row, tenth in enumerate(range(3001), 7):
        time = tenth / 10.0
        values = (
            time, time, 0.0, 2.0, 3.0, 295.0,
            1.0 + 0.2 * math.sin(time / 8.0),
            0.3 + 0.1 * math.cos(time / 10.0),
            0.1 + 0.03 * math.sin(time / 5.0),
            0.0, time - 35.0,
        )
        for column, value in enumerate(values, 1):
            dynamic.cell(row, column, value)
    book.save(path)


def test_replay_package_is_hash_bound_and_keeps_h2_values_sealed(tmp_path):
    workbook = tmp_path / "20200416-Sand03Wind-Final.xlsx"
    report = tmp_path / "PRESLHY-E34-report.docx"
    _workbook(workbook)
    report.write_bytes(b"report fixture")
    source_output = tmp_path / "source-freeze"
    dynamic_output = tmp_path / "dynamic-audit"
    replay_output = tmp_path / "replay"
    FREEZE_TOOL["build_source_window_freeze"](tmp_path, source_output)
    AUDIT_TOOL["build_dynamic_evidence_package"](
        tmp_path, report, dynamic_output,
    )

    record = REPLAY_TOOL["build_replay_input_package"](
        source_output / "source-window-freeze.json",
        dynamic_output / "audit.json",
        trial_id="20200416-Sand03Wind",
        window_ordinal=1,
        output_directory=replay_output,
    )

    assert record["evidence_decision"][
        "dynamic_component_replay_input_allowed"
    ] is True
    assert record["evidence_decision"][
        "obstacle_transport_model_qualification_allowed"
    ] is False
    assert record["sealed_observation_operator"]["values_exported"] is False
    assert record["sealed_observation_operator"]["values_summarized"] is False
    assert record["sealed_observation_operator"]["usable_channels"] == [
        "CH45", "CH55",
    ]
    assert record["inlet_velocity_history"]["interpolation"] == (
        "piecewise_constant"
    )
    assert min(record["inlet_velocity_history"]["u_m_s"]) > 0.0
    assert record["wind_reduction"]["hydrogen_concentration_used"] is False
    assert (replay_output / "dynamic-replay-input-review.png").is_file()
    saved_text = (
        replay_output / "dynamic-replay-input.json"
    ).read_text(encoding="utf-8")
    saved = json.loads(saved_text)
    assert saved["source_freeze"]["sha256"]
    assert "hydrogen_vol_percent" not in saved_text
