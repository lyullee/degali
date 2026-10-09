from pathlib import Path

from docx import Document
import pytest

from tools.build_ijhe_docx import build
from tools.audit_ijhe_docx_draft import audit


@pytest.mark.skipif(
    not Path("outputs/ijhe-main-figure-images-2026-10-09-clean/figure-1.png").is_file(),
    reason="requires locally rendered IJHE figure derivatives",
)
def test_ijhe_docx_builder_preserves_upload_tables_and_placeholders(tmp_path):
    output = tmp_path / "manuscript.docx"
    record = build(
        source=Path("MANUSCRIPT_LH2_VALIDATION.md"),
        output=output,
    )

    assert record["placeholder_state"] == "present"
    document = Document(output)
    assert document.paragraphs[0].text.startswith("Observation-operator-aware")
    assert len(document.tables) >= 2
    assert len(document.inline_shapes) >= 6
    assert any("[AUTHOR TO CONFIRM" in paragraph.text for paragraph in document.paragraphs)
    exported_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "SLABx" not in exported_text
    assert "outputs/" not in exported_text
    assert "tmp/" not in exported_text
    assert "docs/" not in exported_text
    assert "tools/" not in exported_text


def test_ijhe_docx_builder_strict_requires_author_metadata(tmp_path):
    with pytest.raises(ValueError, match="final export requires"):
        build(
            source=Path("MANUSCRIPT_LH2_VALIDATION.md"),
            output=tmp_path / "strict.docx",
            strict=True,
        )


def test_ijhe_docx_builder_strict_rejects_unresolved_declarations(tmp_path):
    source = tmp_path / "source.md"
    source.write_text(
        "# Test manuscript\n\n## Declarations\n\n"
        "Author contributions and competing interests. To be completed by the authors.\n",
        encoding="utf-8",
    )
    with pytest.raises(ValueError, match="completed author declarations"):
        build(
            source=source,
            output=tmp_path / "strict-declarations.docx",
            author_list="A. Author, Institute",
            corresponding="A. Author <a@example.org>",
            strict=True,
        )


@pytest.mark.skipif(
    not Path("outputs/ijhe-manuscript-draft-2026-10-09.docx").is_file(),
    reason="requires the locally rendered IJHE manuscript draft",
)
def test_ijhe_docx_audit_rejects_internal_repository_paths():
    record = audit()
    check = next(item for item in record["checks"] if item["id"] == "docx:internal-paths")
    assert check["status"] == "pass"
