import hashlib
import json
import openpyxl
import pytest

from degali.cli import main
from degali.addons.field_evidence_audit import (
    FIELD_EVIDENCE_AUDIT_SCHEMA,
    FIELD_EVIDENCE_AUDIT_EXECUTION_SCHEMA,
    FieldEvidenceAudit,
    FieldEvidenceCandidate,
    audit_field_evidence,
    field_evidence_audit_record,
    read_field_evidence_audit_json,
    write_field_evidence_audit_json,
)


def _write_complete_channels(root):
    (root / "source.csv").write_text(
        "source_boundary_id,source_rate_kg_s\nsource-a,0.1\n", encoding="utf-8",
    )
    (root / "weather.csv").write_text(
        "weather_id,wind_m_s,temperature\nweather-a,2.0,293\n", encoding="utf-8",
    )
    (root / "obstacle.json").write_text(
        json.dumps({"obstacle_geometry_id": "obstacle-a", "vertices": []}),
        encoding="utf-8",
    )
    (root / "receptors.csv").write_text(
        "sensor,x_downwind_m,y_crosswind_m,height_m,observed_mole_fraction,common_clock_id\n"
        "s1,10,0,1.5,0.04,clock-a\n",
        encoding="utf-8",
    )


def test_audit_detects_channels_but_never_promotes_evidence(tmp_path):
    _write_complete_channels(tmp_path)

    audit = audit_field_evidence(tmp_path)
    assert audit.status == "candidate_complete"
    assert audit.promotion_allowed is False
    assert "cross_file_candidate_unqualified" in audit.gate_codes
    assert "promotion_not_allowed" in audit.gate_codes
    assert all(paths for _, paths in audit.channel_paths)
    assert audit.coherent_candidate_paths == ()
    record = field_evidence_audit_record(audit)
    assert record["schema"] == FIELD_EVIDENCE_AUDIT_SCHEMA
    assert record["promotion_allowed"] is False
    assert record["gate_codes"] == list(audit.gate_codes)
    assert record["coherent_candidate_paths"] == []
    assert record["diagnostic_counts"] == audit.diagnostic_counts
    assert record["collection_requirements"] == audit.collection_requirements
    assert audit.collection_requirements["source_boundary"]["status"] == "detected"
    assert audit.collection_requirements["source_boundary"]["candidate_paths"]
    assert audit.collection_requirements["source_boundary"]["next_action"].startswith(
        "Join a non-negative physical source rate"
    )
    assert all(value == 0 for value in audit.diagnostic_counts.values())
    assert all(candidate.sha256 and len(candidate.sha256) == 64 for candidate in audit.candidates)
    assert all(
        candidate_record["sha256"]
        for candidate_record in record["candidates"]
    )

    source_path = tmp_path / "source.csv"
    original_digest = next(
        candidate.sha256
        for candidate in audit.candidates
        if candidate.path == str(source_path.resolve())
    )
    source_path.write_text(
        "source_boundary_id,source_rate_kg_s\nsource-a,0.2\n", encoding="utf-8",
    )
    refreshed = audit_field_evidence(tmp_path)
    refreshed_digest = next(
        candidate.sha256
        for candidate in refreshed.candidates
        if candidate.path == str(source_path.resolve())
    )
    assert refreshed_digest != original_digest


def test_pool_radius_observations_do_not_count_as_receptor_validation(tmp_path):
    (tmp_path / "radius.csv").write_text(
        "source_record_id,surface,trial,time_s,observed_radius_m\n"
        "JUEL,water,3,2,0.4\n",
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)
    assert audit.status == "withheld"
    assert "channels_missing" in audit.gate_codes
    assert all(not paths for _, paths in audit.channel_paths)


def test_unidentified_release_rate_is_not_a_source_boundary(tmp_path):
    (tmp_path / "literature-inventory.json").write_text(
        json.dumps({
            "release_rate_kg_s": 0.5,
            "source_rate_kg_s": 0.5,
            "mass_rate_reported_kg_s": 0.5,
            "source_temperature_K": 20.0,
        }),
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)
    assert audit.status == "withheld"
    assert not dict(audit.channel_paths)["source_boundary"]
    assert any(
        "source-boundary identifier" in note
        for candidate in audit.candidates
        for note in candidate.notes
    )
    assert audit.diagnostic_counts["release_rate_without_source_identity"] == 1


def test_operating_pressure_history_is_not_an_atmospheric_source(tmp_path):
    (tmp_path / "vent-historian.csv").write_text(
        "timestamp,tank_pressure,tank_temperature,liquid_level\n"
        "0,320000,23.5,0.82\n"
        "1,319000,23.6,0.81\n",
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    assert not dict(audit.channel_paths)["source_boundary"]
    assert audit.diagnostic_counts["operating_history_without_source_boundary"] == 1
    assert any(
        "not a dispersion source" in note
        for candidate in audit.candidates
        for note in candidate.notes
    )


def test_instrument_tag_history_is_not_an_atmospheric_source(tmp_path):
    (tmp_path / "tag-history.csv").write_text(
        "timestamp,PT1101,TT1106,LT1101\n"
        "0,3.2e5,23.5,0.82\n"
        "1,3.19e5,23.6,0.81\n",
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    assert not dict(audit.channel_paths)["source_boundary"]
    assert audit.diagnostic_counts["instrument_history_without_source_boundary"] == 1


def test_header_only_structured_files_do_not_create_readiness_channels(tmp_path):
    (tmp_path / "empty.csv").write_text(
        "source_boundary_id,mass_flow_kg_s,weather_id,wind_speed_m_s,"
        "sensor_id,x_m,y_m,z_m,observed_mole_fraction,common_clock_id\n",
        encoding="utf-8",
    )
    (tmp_path / "empty.json").write_text("[]", encoding="utf-8")

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    assert all(not paths for _, paths in audit.channel_paths)
    assert any(
        "contains no data rows" in note
        for candidate in audit.candidates
        for note in candidate.notes
    )
    assert audit.diagnostic_counts["empty_structured_file"] == 2


def test_blank_structured_records_do_not_create_readiness_channels(tmp_path):
    (tmp_path / "blank.csv").write_text(
        "source_boundary_id,mass_flow_kg_s\n,\n", encoding="utf-8",
    )
    (tmp_path / "blank.json").write_text(
        json.dumps({"observations": [{}]}), encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    assert all(not paths for _, paths in audit.channel_paths)
    assert sum(
        "contains no data rows" in note
        for candidate in audit.candidates
        for note in candidate.notes
    ) == 2


@pytest.mark.parametrize(
    ("name", "contents"),
    (
        (
            "duplicate-header.csv",
            "source_boundary_id,mass_flow_kg_s,mass_flow_kg_s\nsource-a,0.1,0.2\n",
        ),
        (
            "extra-value.csv",
            "source_boundary_id,mass_flow_kg_s\nsource-a,0.1,unexpected\n",
        ),
        (
            "empty-header.csv",
            "source_boundary_id,,mass_flow_kg_s\nsource-a,unused,0.1\n",
        ),
    ),
)
def test_malformed_csv_structure_is_not_promoted(tmp_path, name, contents):
    (tmp_path / name).write_text(contents, encoding="utf-8")

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    assert all(not paths for _, paths in audit.channel_paths)
    candidate = audit.candidates[0]
    assert candidate.detected_channels == ()
    assert any("unreadable structured file: ValueError" in note for note in candidate.notes)


def test_large_csv_is_stream_validated_but_channel_sample_is_bounded(tmp_path):
    path = tmp_path / "large.csv"
    rows = ["source_boundary_id,mass_flow_kg_s\n"]
    rows.extend("source-a,0.1\n" for _ in range(250))
    path.write_text("".join(rows), encoding="utf-8")

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "partial"
    assert "channels_partial" in audit.gate_codes
    candidate = audit.candidates[0]
    assert candidate.row_count == 250
    assert "source_boundary" in candidate.detected_channels
    assert any("first 200 data rows" in note for note in candidate.notes)


@pytest.mark.parametrize(
    "payload",
    (
        '{"source_id":"source-a","source_id":"source-b"}',
        '{"source_id":"source-a","source-id":"source-b"}',
    ),
)
def test_duplicate_or_normalized_json_keys_are_not_promoted(tmp_path, payload):
    (tmp_path / "duplicate.json").write_text(payload, encoding="utf-8")

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    assert all(not paths for _, paths in audit.channel_paths)
    candidate = audit.candidates[0]
    assert candidate.detected_channels == ()
    assert any("unreadable structured file: ValueError" in note for note in candidate.notes)


def test_placeholder_values_do_not_create_readiness_channels(tmp_path):
    (tmp_path / "placeholder.csv").write_text(
        "source_boundary_id,mass_flow_kg_s,weather_id,wind_speed_m_s,"
        "sensor_id,x_m,y_m,z_m,observed_mole_fraction,common_clock_id\n"
        "source-a,n/a,weather-a,calm,H2-01,unknown,0,0,not-recorded,n/a\n",
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    assert all(not paths for _, paths in audit.channel_paths)


def test_channels_require_fields_to_cooccur_in_one_record(tmp_path):
    (tmp_path / "split-source.csv").write_text(
        "source_boundary_id,mass_flow_kg_s\nsource-a,\n,0.1\n",
        encoding="utf-8",
    )
    (tmp_path / "split-receptor.csv").write_text(
        "sensor_id,x_m,y_m,z_m,observed_mole_fraction,common_clock_id\n"
        "S1,10,,1,0.04,clock-a\n"
        "S2,,0,1,0.03,clock-a\n",
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)

    paths = dict(audit.channel_paths)
    assert not paths["source_boundary"]
    assert not paths["receptor_observations"]
    assert paths["common_clock"]


@pytest.mark.parametrize("clock_column", [
    "synchronization_id", "sync_id", "time_sync_id", "clock_sync_id",
])
def test_common_clock_aliases_are_detected_only_with_nonblank_values(tmp_path, clock_column):
    path = tmp_path / "clock-alias.csv"
    path.write_text(
        f"sensor_id,timestamp,{clock_column}\nS1,2026-01-01T00:00:00Z,clock-a\n",
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)

    assert dict(audit.channel_paths)["common_clock"] == (str(path.resolve()),)
    assert audit.collection_requirements["common_clock"]["status"] == "detected"
    assert "sync_id" in audit.collection_requirements["common_clock"][
        "required_field_groups"
    ][0]

    path.write_text(
        f"sensor_id,timestamp,{clock_column}\nS1,2026-01-01T00:00:00Z,\n",
        encoding="utf-8",
    )
    refreshed = audit_field_evidence(tmp_path)
    assert not dict(refreshed.channel_paths)["common_clock"]


@pytest.mark.parametrize("metadata_column", [
    "synchronization_method", "time_alignment_id", "common_time_id", "clock_offset_s",
])
def test_timestamp_synchronization_metadata_is_a_common_clock_candidate(tmp_path, metadata_column):
    path = tmp_path / "clock-metadata.csv"
    path.write_text(
        f"sensor_id,timestamp,{metadata_column}\nS1,2026-01-01T00:00:00Z,aligned\n",
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)

    assert dict(audit.channel_paths)["common_clock"] == (str(path.resolve()),)

    path.write_text(
        f"sensor_id,timestamp,{metadata_column}\nS1,2026-01-01T00:00:00Z,\n",
        encoding="utf-8",
    )
    refreshed = audit_field_evidence(tmp_path)
    assert not dict(refreshed.channel_paths)["common_clock"]


def test_common_external_aliases_are_detected_without_cross_file_inference(tmp_path):
    cases = (
        (
            "source",
            "source_bound,source_rate_kg_s\nboundary-a,0.1\n",
            "source_boundary",
            "source_bound",
        ),
        (
            "weather",
            "weather_id,wind_ref_m_s\nweather-a,2.0\n",
            "weather",
            "wind_ref_m_s",
        ),
        (
            "receptor",
            "sensor,x_m,y_m,z_m,mean_c_pct\nS1,1.0,0.0,1.5,0.02\n",
            "receptor_observations",
            "mean_c_pct",
        ),
    )
    for name, content, channel, alias in cases:
        case_root = tmp_path / name
        case_root.mkdir()
        path = case_root / "evidence.csv"
        path.write_text(content, encoding="utf-8")

        audit = audit_field_evidence(case_root)

        assert dict(audit.channel_paths)[channel] == (str(path.resolve()),)
        assert alias in audit.collection_requirements[channel][
            "required_field_groups"
        ][0 if channel == "source_boundary" else -1]


def test_json_event_metadata_is_checked_with_each_observation_record(tmp_path):
    (tmp_path / "event.json").write_text(
        json.dumps({
            "source_boundary_id": "source-a",
            "mass_flow_kg_s": 0.1,
            "weather_id": "weather-a",
            "wind_speed_m_s": 2.0,
            "obstacle_geometry_id": "obstacle-a",
            "observations": [{
                "sensor_id": "H2-01",
                "x_m": 1.0,
                "y_m": 0.0,
                "z_m": 0.5,
                "observed_mole_fraction": 0.04,
                "common_clock_id": "clock-a",
            }],
        }),
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "candidate_complete"
    assert "coherent_candidate_unqualified" in audit.gate_codes
    assert audit.max_files == 20000
    assert audit.promotion_allowed is False
    assert audit.coherent_candidate_paths == (str((tmp_path / "event.json").resolve()),)


def test_threshold_sensor_summary_is_retained_as_nonquantitative_note(tmp_path):
    path = tmp_path / "published_sensor_summary.csv"
    path.write_text(
        "sensor_id,radius_m,angle_deg,height_m,first_above_4pct_s,final_below_4pct_s\n"
        "H49,2.5,-45,4,8,210\n",
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)
    assert audit.status == "withheld"
    assert not dict(audit.channel_paths)["receptor_observations"]
    assert any(
        "not a validation observation" in note
        for candidate in audit.candidates
        for note in candidate.notes
    )


def test_observed_time_mean_volume_percent_is_a_receptor_observation(tmp_path):
    (tmp_path / "matched-sensors.csv").write_text(
        "test,sensor,x_downwind_m,y_crosswind_m,height_m,"
        "averaging_time_s,observed_time_mean_vol_pct\n"
        "4,OC_01,21.2,-21.2,1.0,275,0.1\n",
        encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)

    assert dict(audit.channel_paths)["receptor_observations"]
    assert audit.status == "partial"


def test_unidentified_weather_fields_are_not_promoted(tmp_path):
    (tmp_path / "weather.csv").write_text(
        "wind_m_s,mean_wind_from_deg,temperature\n2.0,180,293.15\n", encoding="utf-8",
    )
    audit = audit_field_evidence(tmp_path)
    assert audit.status == "withheld"
    assert not dict(audit.channel_paths)["weather"]
    assert any("weather identifier" in note for c in audit.candidates for note in c.notes)
    assert audit.diagnostic_counts["weather_without_identity"] == 1


def test_audit_reports_partial_channels_and_cli_record(tmp_path, capsys):
    (tmp_path / "clock.csv").write_text(
        "sensor,common_clock_id\ns1,clock-a\n", encoding="utf-8",
    )

    audit = audit_field_evidence(tmp_path)
    assert audit.status == "partial"
    assert dict(audit.channel_paths)["common_clock"]
    assert main(["field-audit", str(tmp_path)]) == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["field_evidence_audit"]["status"] == "partial"
    requirements = payload["field_evidence_audit"]["collection_requirements"]
    assert requirements["common_clock"]["status"] == "detected"
    assert requirements["source_boundary"]["status"] == "missing"
    assert "Export the event's atmospheric source-boundary ID" in requirements[
        "source_boundary"
    ]["next_action"]
    output = tmp_path / "cli-audit.json"
    assert main(["field-audit", str(tmp_path), "--output", str(output)]) == 0
    assert read_field_evidence_audit_json(output).status == "partial"
    capsys.readouterr()
    assert main(["field-audit-verify", str(output)]) == 0
    verification = json.loads(capsys.readouterr().out)
    assert verification["schema"] == "degali.field-evidence-audit-verification.v1"
    assert verification["verification"]["candidate_files_verified"] is True
    assert main(["field-audit-verify", str(output), "--require-complete"]) == 2
    capsys.readouterr()


def test_audit_json_writer_round_trips_and_rejects_mutation(tmp_path):
    _write_complete_channels(tmp_path)
    audit = audit_field_evidence(tmp_path)
    output = tmp_path / "audit.json"

    parsed = write_field_evidence_audit_json(audit, output)

    assert parsed.as_record() == audit.as_record()
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["schema"] == FIELD_EVIDENCE_AUDIT_EXECUTION_SCHEMA
    assert payload["input"]["max_files"] == audit.max_files
    assert read_field_evidence_audit_json(output).as_record() == audit.as_record()
    with pytest.raises(FileExistsError, match="overwrite"):
        write_field_evidence_audit_json(audit, output)

    payload["input"]["max_files"] += 1
    output.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValueError, match="input.max_files"):
        read_field_evidence_audit_json(output)
    payload["input"]["max_files"] = audit.max_files
    output.write_text(json.dumps(payload), encoding="utf-8")

    payload["field_evidence_audit"]["diagnostic_counts"][
        "release_rate_without_source_identity"
    ] += 1
    output.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValueError, match="diagnostic_counts must match candidate notes"):
        read_field_evidence_audit_json(output)
    payload["field_evidence_audit"]["diagnostic_counts"] = audit.diagnostic_counts
    output.write_text(json.dumps(payload), encoding="utf-8")

    payload["field_evidence_audit"]["collection_requirements"][
        "source_boundary"
    ]["status"] = "qualified"
    output.write_text(json.dumps(payload), encoding="utf-8")
    with pytest.raises(ValueError, match="collection_requirements"):
        read_field_evidence_audit_json(output)
    payload["field_evidence_audit"]["collection_requirements"] = (
        audit.collection_requirements
    )
    output.write_text(json.dumps(payload), encoding="utf-8")

    source = tmp_path / "source.csv"
    source.write_text(
        "source_boundary_id,mass_flow_kg_s\nsource-a,0.3\n", encoding="utf-8",
    )
    with pytest.raises(ValueError, match="changed since its provenance"):
        read_field_evidence_audit_json(output)


def test_audit_artifact_without_optional_collection_plan_remains_readable(tmp_path):
    _write_complete_channels(tmp_path)
    audit = audit_field_evidence(tmp_path)
    output = tmp_path / "legacy-audit.json"
    write_field_evidence_audit_json(audit, output)

    payload = json.loads(output.read_text(encoding="utf-8"))
    del payload["field_evidence_audit"]["collection_requirements"]
    output.write_text(json.dumps(payload), encoding="utf-8")

    parsed = read_field_evidence_audit_json(output)
    assert parsed.status == "candidate_complete"
    assert parsed.collection_requirements == audit.collection_requirements


def test_audit_artifact_with_legacy_clock_alias_plan_remains_readable(tmp_path):
    _write_complete_channels(tmp_path)
    audit = audit_field_evidence(tmp_path)
    output = tmp_path / "legacy-clock-plan.json"
    write_field_evidence_audit_json(audit, output)

    payload = json.loads(output.read_text(encoding="utf-8"))
    payload["field_evidence_audit"]["collection_requirements"]["common_clock"][
        "required_field_groups"
    ] = [["common_clock_id", "clock_id", "synchronized_clock_id"]]
    output.write_text(json.dumps(payload), encoding="utf-8")

    parsed = read_field_evidence_audit_json(output)
    assert parsed.status == "candidate_complete"


def test_audit_artifact_binds_channel_paths_to_candidate_channels(tmp_path):
    _write_complete_channels(tmp_path)
    output = tmp_path / "audit.json"
    audit = audit_field_evidence(tmp_path)
    write_field_evidence_audit_json(audit, output)

    payload = json.loads(output.read_text(encoding="utf-8"))
    channels = payload["field_evidence_audit"]["channel_paths"]
    channels["source_boundary"] = []
    output.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(
        ValueError, match="channel_paths must match candidate detected_channels",
    ):
        read_field_evidence_audit_json(output)


def test_audit_artifact_rejects_candidate_path_outside_root(tmp_path):
    root = tmp_path / "root"
    root.mkdir()
    outside = tmp_path / "outside.csv"
    outside.write_text("source_boundary_id,mass_flow_kg_s\nsource-a,0.1\n", encoding="utf-8")
    candidate = FieldEvidenceCandidate(
        path=str(outside),
        file_type="csv",
        row_count=1,
        fields=("source_boundary_id", "mass_flow_kg_s"),
        detected_channels=("source_boundary",),
        sha256=hashlib.sha256(outside.read_bytes()).hexdigest(),
    )
    channels = (
        "source_boundary", "weather", "obstacle_geometry",
        "receptor_observations", "common_clock",
    )
    audit = FieldEvidenceAudit(
        root=str(root),
        scanned_files=1,
        candidates=(candidate,),
        status="partial",
        channel_paths=tuple(
            (channel, (candidate.path,) if channel == "source_boundary" else ())
            for channel in channels
        ),
        reasons=("outside-root test",),
    )

    with pytest.raises(ValueError, match="escapes the audit root"):
        write_field_evidence_audit_json(audit, tmp_path / "outside-root-audit.json")


def test_audit_rejects_file_limit(tmp_path):
    (tmp_path / "a_unrelated.csv").write_text("unrelated,value\nnone,0\n", encoding="utf-8")
    (tmp_path / "z_source.csv").write_text(
        "source_boundary_id,mass_flow_kg_s\nsource-a,0.1\n", encoding="utf-8",
    )
    audit = audit_field_evidence(tmp_path, max_files=1)
    assert audit.max_files == 1
    assert audit.scan_complete is False
    assert audit.status == "withheld"
    assert any("scan stopped" in reason for reason in audit.reasons)
    assert "scan_incomplete" in audit.gate_codes

    try:
        audit_field_evidence(tmp_path, max_files=0)
    except ValueError as error:
        assert "max_files" in str(error)
    else:
        raise AssertionError("expected max_files validation")


def test_audit_samples_xlsx_headers_without_promoting_workbook_identity(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "event"
    sheet.append([
        "source_boundary_id", "mass_flow_kg_s", "weather_id", "wind_speed_m_s",
        "obstacle_geometry_id", "sensor_id", "x_m", "y_m", "z_m",
        "observed_mole_fraction", "common_clock_id",
    ])
    sheet.append([
        "source-a", 0.1, "weather-a", 2.0, "obstacle-a", "H2-01",
        1.0, 0.0, 0.5, 0.04, "clock-a",
    ])
    path = tmp_path / "field-event.xlsx"
    workbook.save(path)
    workbook.close()

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "candidate_complete"
    assert audit.promotion_allowed is False
    candidate = audit.candidates[0]
    assert candidate.file_type == "xlsx"
    assert candidate.row_count is None
    assert any("sampled" in note for note in candidate.notes)


def test_audit_preserves_unicode_xlsx_headers_without_calling_them_corrupt(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "현장기록"
    sheet.append(["회차", "timestamp", "탱크압력", "온도"])
    sheet.append([1, "00:00", 0.31, -33.8])
    path = tmp_path / "unicode-fields.xlsx"
    workbook.save(path)
    workbook.close()

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    candidate = audit.candidates[0]
    assert candidate.detected_channels == ()
    assert "회차" in candidate.fields
    assert "탱크압력" in candidate.fields
    assert not any("unreadable structured file" in note for note in candidate.notes)


def test_audit_keeps_valid_xlsx_sheets_when_an_auxiliary_sheet_is_ambiguous(tmp_path):
    workbook = openpyxl.Workbook()
    valid = workbook.active
    valid.title = "측정"
    valid.append(["source_boundary_id", "mass_flow_kg_s"])
    valid.append(["source-a", 0.1])
    auxiliary = workbook.create_sheet("요약")
    auxiliary.append(["제목", None, "설명"])
    auxiliary.append(["값", None, "메모"])
    path = tmp_path / "mixed-sheets.xlsx"
    workbook.save(path)
    workbook.close()

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "partial"
    assert audit.channel_paths[0][0] == "source_boundary"
    assert audit.channel_paths[0][1] == (str(path.resolve()),)
    candidate = audit.candidates[0]
    assert "source_boundary" in candidate.detected_channels
    assert any("worksheet '요약' skipped" in note for note in candidate.notes)
    assert not any("unreadable structured file" in note for note in candidate.notes)


def test_header_only_xlsx_does_not_create_readiness_channels(tmp_path):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append([
        "source_boundary_id", "mass_flow_kg_s", "weather_id", "wind_speed_m_s",
        "obstacle_geometry_id", "sensor_id", "x_m", "y_m", "z_m",
        "observed_mole_fraction", "common_clock_id",
    ])
    path = tmp_path / "header-only.xlsx"
    workbook.save(path)
    workbook.close()

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    assert all(not paths for _, paths in audit.channel_paths)
    candidate = audit.candidates[0]
    assert candidate.row_count is None
    assert any("contains no data rows" in note for note in candidate.notes)


@pytest.mark.parametrize(
    "headers",
    [
        ["source_boundary_id", "source_boundary_id", "mass_flow_kg_s"],
        ["source_boundary_id", "", "mass_flow_kg_s"],
    ],
)
def test_audit_rejects_ambiguous_xlsx_headers(tmp_path, headers):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(headers)
    sheet.append(["source-a", "source-a", 0.1])
    path = tmp_path / "ambiguous.xlsx"
    workbook.save(path)
    workbook.close()

    audit = audit_field_evidence(tmp_path)

    assert audit.status == "withheld"
    candidate = audit.candidates[0]
    assert candidate.detected_channels == ()
    assert candidate.notes == ("unreadable structured file: ValueError",)


def test_audit_contract_rejects_malformed_candidate_and_partial_state():
    with pytest.raises(ValueError, match="detected_channels must be unique"):
        FieldEvidenceCandidate(
            path="candidate.csv",
            file_type="csv",
            row_count=1,
            fields=("weather_id",),
            detected_channels=("weather", "weather"),
        )

    channels = (
        "source_boundary", "weather", "obstacle_geometry",
        "receptor_observations", "common_clock",
    )
    with pytest.raises(ValueError, match="partial audit must contain"):
        FieldEvidenceAudit(
            root="C:/evidence",
            scanned_files=0,
            candidates=(),
            status="partial",
            channel_paths=tuple((channel, ()) for channel in channels),
            reasons=("malformed",),
            scan_complete=True,
        )

    with pytest.raises(ValueError, match="unsupported code"):
        FieldEvidenceAudit(
            root="C:/evidence",
            scanned_files=0,
            candidates=(),
            status="withheld",
            channel_paths=tuple((channel, ()) for channel in channels),
            reasons=("malformed",),
            gate_codes=("not-a-code",),
        )


def test_audit_rejects_coherent_path_without_all_channels():
    candidate = FieldEvidenceCandidate(
        path="candidate.csv",
        file_type="csv",
        row_count=1,
        fields=("source_boundary_id",),
        detected_channels=("source_boundary",),
    )
    channels = (
        "source_boundary", "weather", "obstacle_geometry",
        "receptor_observations", "common_clock",
    )
    with pytest.raises(ValueError, match="every required channel"):
        FieldEvidenceAudit(
            root="C:/evidence",
            scanned_files=1,
            candidates=(candidate,),
            status="partial",
            channel_paths=tuple(
                (channel, ("candidate.csv",) if channel == "source_boundary" else ())
                for channel in channels
            ),
            reasons=("malformed",),
            coherent_candidate_paths=("candidate.csv",),
        )
