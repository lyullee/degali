from openpyxl import Workbook

from degali.validation.hytunnel_ps import (
    _condition_from_title,
    read_hytunnel_concentration_workbook,
    read_hytunnel_flow_workbook,
)


def test_component_readers_are_exposed_by_validation_public_api():
    import degali.validation as validation

    assert validation.read_hytunnel_concentration_workbook is read_hytunnel_concentration_workbook
    assert callable(validation.read_usn_carpark_mat)


def _matrix_sheet(sheet, *, flow=False):
    if flow:
        sheet["B1"] = "average value m/s (1 min duration)"
        sheet["J1"] = "max value m/s (1 min duration)"
        sheet["R1"] = "min value m/s (1 min duration)"
        for column in (2, 10, 18):
            sheet.cell(3, column, "y / mm")
            sheet.cell(3, column + 1, -500)
            sheet.cell(3, column + 2, 0)
            sheet.cell(4, column, "x / mm")
            sheet.cell(5, column, 250)
            sheet.cell(5, column + 1, 1.0)
            sheet.cell(5, column + 2, 2.0)
        return
    sheet["K3"] = "average value"
    sheet["V3"] = "maximum value"
    sheet["AH3"] = "minimum value"
    for column in (10, 22, 34):
        sheet.cell(6, column, "x/ mm")
        sheet.cell(6, column + 1, -175)
        sheet.cell(6, column + 2, 0)
        sheet.cell(7, column, 125)
        sheet.cell(7, column + 1, 1.0)
        sheet.cell(7, column + 2, 2.0)


def test_hytunnel_concentration_reader_preserves_three_statistics(tmp_path):
    path = tmp_path / "concentration.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "1 mm 5g H2 W = 3.5 cross-flow"
    _matrix_sheet(sheet)
    workbook.save(path)

    dataset = read_hytunnel_concentration_workbook(path)
    assert len(dataset.configurations) == 1
    assert dataset.configurations[0].flow_direction == "cross-flow"
    assert dataset.configurations[0].is_direct_hydrogen
    assert {point.statistic for point in dataset.points} == {
        "average", "maximum", "minimum",
    }
    assert len(dataset.points) == 6
    assert not dataset.quantitative_lh2_validation_allowed


def test_hytunnel_flow_reader_reads_signed_velocity(tmp_path):
    path = tmp_path / "flow.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "3.5 m pro s counter-flow"
    _matrix_sheet(sheet, flow=True)
    sheet["C5"] = -1.0
    workbook.save(path)

    dataset = read_hytunnel_flow_workbook(path)
    assert len(dataset.configuration_ids) == 1
    assert len(dataset.points) == 6
    assert min(point.velocity_x_m_s for point in dataset.points) == -1.0


def test_condition_parser_accepts_published_missing_w_typo():
    assert _condition_from_title("1 mm 5g H2 = 5 counter-flow") == (
        1.0, 5.0, 5.0, "counter-flow", "H2"
    )
