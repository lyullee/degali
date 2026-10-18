from pathlib import Path
import json
import runpy

import openpyxl
import pytest

from degali.addons.field_evidence_manifest import (
    field_evidence_manifest_from_mapping,
)


ROOT = Path(__file__).resolve().parents[1]
TOOL = runpy.run_path(ROOT / "tools/audit_preslhy_e34_dynamic_evidence.py")


def _workbook(path: Path, *, defective_wind: bool = False) -> None:
    book = openpyxl.Workbook()
    pool = book.active
    pool.title = "TE-WG"
    pool.cell(7, 1, "X_Value")
    pool.cell(7, 2, "TA000-00")
    pool.cell(7, 31, "Sync. Time [s]")
    pool.cell(7, 33, "m(LH2) [g]")
    for row, time in enumerate((0.0, 1.0, 2.0), 8):
        pool.cell(row, 1, time)
        pool.cell(row, 2, 293.0 - time)
        pool.cell(row, 31, time)
        pool.cell(row, 33, 1000.0 - 10.0 * time)

    dynamic = book.create_sheet("cH2-Wind")
    headers = (
        "Orig.Time [s]", "Sync. Time [s]", "CH35", "CH45", "CH55",
        "Young_T", "Young_U", "Young_V", "Young_W", "Trigger",
        "cH2 Time [s]",
    )
    for column, label in enumerate(headers, 1):
        dynamic.cell(6, column, label)
        dynamic.cell(5, column, 0.1)
    dynamic["C5"] = "defective"
    dynamic["C5"].fill = openpyxl.styles.PatternFill(
        fill_type="solid", fgColor="FFFFCCFF",
    )
    if defective_wind:
        for cell in ("G5", "H5", "I5"):
            dynamic[cell].fill = openpyxl.styles.PatternFill(
                fill_type="solid", fgColor="FFFFCCFF",
            )
    for row, time in enumerate((0.0, 0.1, 0.2), 7):
        values = (
            time, time - 0.4, 0.4, 1.0 + time, 2.0 + time,
            295.0, 1.0, 2.0, 0.2, 0.0, time - 35.4,
        )
        for column, value in enumerate(values, 1):
            dynamic.cell(row, column, value)
    book.save(path)


def test_dynamic_evidence_manifest_is_conditional_and_fail_closed(tmp_path):
    workbook = tmp_path / "20200416-Sand03Wind-Final.xlsx"
    report = tmp_path / "PRESLHY-E34-report.docx"
    _workbook(workbook)
    report.write_bytes(b"report fixture")

    record = TOOL["audit_dynamic_trial"](
        workbook, report, audit_root=tmp_path,
    )

    manifest = record["field_evidence_manifest"]
    decision = record["evidence_decision"]
    assert manifest["event_id"] == "20200416-Sand03Wind"
    assert manifest["evidence_readiness"]["status"] == "conditional"
    assert manifest["promotion_allowed"] is False
    assert len(manifest["channel_artifacts"]) == 5
    assert decision["dynamic_component_replay_allowed"] is True
    assert decision["obstacle_transport_model_qualification_allowed"] is False
    assert record["dynamic_history"]["usable_hydrogen_channels"] == [
        "CH45", "CH55",
    ]
    assert record["dynamic_history"]["sampling_frequency_hz"] == pytest.approx(10.0)


def test_workbook_flagged_wind_withholds_dynamic_replay(tmp_path):
    workbook = tmp_path / "20200415-Sand02Wind-Final.xlsx"
    report = tmp_path / "PRESLHY-E34-report.docx"
    _workbook(workbook, defective_wind=True)
    report.write_bytes(b"report fixture")

    record = TOOL["audit_dynamic_trial"](
        workbook, report, audit_root=tmp_path,
    )

    assert record["evidence_decision"]["status"] == "withheld"
    assert record["evidence_decision"]["dynamic_component_replay_allowed"] is False
    assert "workbook_flagged_defective_sonic_wind_channels" in record[
        "evidence_decision"
    ]["reasons"]


def test_saved_dynamic_manifest_round_trips_and_verifies_raw_files(tmp_path):
    workbook = tmp_path / "20200416-Sand03Wind-Final.xlsx"
    report = tmp_path / "PRESLHY-E34-report.docx"
    output = tmp_path / "package"
    _workbook(workbook)
    report.write_bytes(b"report fixture")

    record = TOOL["build_dynamic_evidence_package"](
        tmp_path, report, output,
    )

    saved = json.loads(
        (output / record["manifest_files"][0]).read_text(encoding="utf-8")
    )
    manifest = field_evidence_manifest_from_mapping(saved)
    manifest.verify_files()
    assert manifest.event_id == "20200416-Sand03Wind"
    assert manifest.readiness_record()["status"] == "conditional"
