"""PRESLHY table parsing and near-field temperature reduction."""

import openpyxl
import pytest

from degali.validation.preslhy import (
    _table_a3_positions,
    compensate_pipe_temperature_c,
    read_nearfield_temperatures,
    read_pipe_source,
    type_t_emf_mv,
    type_t_temperature_c,
)


def test_table_a3_keeps_unicode_negative_and_long_serials():
    text = (
        "02EC40 M3S2 1.78 ‐ 1.00 0.50 TC2333 "
        "10DC100 M2S6 1.78 0.00 0.00 TC2324"
    )
    positions = _table_a3_positions(text)
    assert positions["02EC40"] == (1.78, -1.0, 0.5)
    assert positions["10DC100"] == (1.78, 0.0, 0.0)


def test_nist_type_t_anchor_and_inverse():
    emf = type_t_emf_mv(-253.0)
    # Published coefficients are decimal constants; evaluation order differs
    # by about 1e-11 mV, far below their tabulated 0.001 mV resolution.
    assert emf == pytest.approx(-6.198375887745304, abs=1e-10)
    assert type_t_temperature_c(emf) == pytest.approx(-253.0, abs=1e-10)


def test_preslhy_pipe_gain_correction_maps_tc1_minimum_to_minus_253_c():
    corrected = compensate_pipe_temperature_c(-7023.476, -7023.476)
    assert corrected == pytest.approx(-253.0, abs=1e-10)
    # Trial 10's TC3 median is slightly warmer than the inlet reference.
    nozzle = compensate_pipe_temperature_c(-7019.416, -7023.476)
    assert -253.0 < nozzle < -252.0


def test_pipe_source_reader_uses_raw_tc3_and_frozen_window(tmp_path):
    path = tmp_path / "trial_23_test.xlsx"
    book = openpyxl.Workbook()
    flex = book.active
    flex.title = "Flexlogger"
    flex.append([
        "MFM1_Mass_Flow_Rate", "PT1_Pipe_Pressure",
        "PT2_Nozzle_Pressure", "TC1_MFM_In",
        "TC3_Release_Nozzle_Flow",
    ])
    flex.append([100.0, 2.6, 2.2, -7023.476, -7019.416])
    flex.append([120.0, 2.7, 2.3, -7000.0, -7000.0])
    flex.append([140.0, 2.8, 2.4, -7010.0, -6980.0])
    flow = book.create_sheet("Flowmeter")
    flow.append([
        "FlMassFlowRategsecR0247", "FlDriveGainR0291",
        "FlCalcDensitykgm3",
    ])
    flow.append([101.0, 99.0, 35.0])
    flow.append([121.0, 100.0, 36.0])
    flow.append([141.0, 100.0, 37.0])
    book.save(path)

    source = read_pipe_source(path, (0, 3))
    assert source.trial == 23
    assert source.samples == 3
    assert source.pt1_median_barg == pytest.approx(2.7)
    assert source.pt2_median_barg == pytest.approx(2.3)
    assert source.flex_mass_flow_mean_g_s == pytest.approx(120.0)
    assert source.coriolis_mass_flow_median_g_s == pytest.approx(121.0)
    assert source.drive_gain_median_pct == pytest.approx(100.0)
    assert source.calculated_density_median_kg_m3 == pytest.approx(36.0)
    assert source.nozzle_temperature_median_k == pytest.approx(
        compensate_pipe_temperature_c(-7000.0, -7023.476) + 273.15
    )


def test_temperature_reduction_uses_flow_window_and_converted_channels(tmp_path):
    path = tmp_path / "trial_10_test.xlsx"
    book = openpyxl.Workbook()
    flow = book.active
    flow.title = "Flowmeter"
    flow.append(["FlTime", "FlMassFlowRategsecR0247"])
    for index, value in enumerate([0, 0, 10, 10, 10, 10, 10, 0, 0]):
        flow.append([index, value])
    gas = book.create_sheet("Xensor")
    gas.append(["X2019_09_04_02EC25Time", "X2019_09_04_02EC25Output"])
    for index in range(27):
        gas.append([index, 1.0])
    flex = book.create_sheet("Flexlogger")
    flex.append([
        "time", "Centreline_178_TC2328", "Centreline_178_TC2328C",
        "TC1_MFM_InC",
    ])
    for index, value in enumerate([20, 20, -10, -20, -30, -40, -50, 20, 20]):
        flex.append([index, -100.0, value, -253.0])
    book.save(path)

    window, readings = read_nearfield_temperatures(path)
    assert window == (2, 7)
    assert len(readings) == 1
    reading = readings[0]
    assert reading.serial == "TC2328"
    assert (reading.x, reading.y, reading.z) == (1.78, 0.0, 0.5)
    assert reading.minimum_c == -50.0
    assert reading.median_c == -30.0
    assert reading.samples == 5
