"""Audit the locally generated IJHE DOCX/PDF draft without promoting it."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
SCHEMA = "degali.ijhe-docx-draft-audit.v1"
DOCX = ROOT / "outputs" / "ijhe-manuscript-draft-2026-10-09.docx"
PDF = ROOT / "outputs" / "ijhe-manuscript-draft-2026-10-09.pdf"


def audit() -> dict[str, Any]:
    checks: list[dict[str, Any]] = []

    if not DOCX.is_file():
        checks.append({"id": "docx:present", "status": "warning", "detail": "draft DOCX is absent"})
        return {
            "schema": SCHEMA,
            "status": "not_available",
            "final_upload_allowed": False,
            "checks": checks,
        }

    document = Document(DOCX)
    title = document.paragraphs[0].text if document.paragraphs else ""
    document_text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    internal_path_markers = tuple(
        marker for marker in ("docs/", "tools/", "outputs/", "tmp/") if marker in document_text
    )
    placeholder_count = sum(
        marker in paragraph.text
        for paragraph in document.paragraphs
        for marker in ("[AUTHOR TO CONFIRM", "[CORRESPONDING AUTHOR TO CONFIRM")
    )
    checks.extend(
        [
            {
                "id": "docx:present",
                "status": "pass",
                "detail": f"{DOCX.name} exists ({DOCX.stat().st_size} bytes)",
            },
            {
                "id": "docx:title",
                "status": "pass" if title.startswith("Observation-operator-aware") else "fail",
                "detail": title,
            },
            {
                "id": "docx:tables",
                "status": "pass" if len(document.tables) >= 2 else "fail",
                "detail": f"{len(document.tables)} tables present",
            },
            {
                "id": "docx:author-placeholders",
                "status": "warning" if placeholder_count else "pass",
                "detail": f"{placeholder_count} unresolved author placeholder paragraphs",
            },
            {
                "id": "docx:internal-paths",
                "status": "fail" if internal_path_markers else "pass",
                "detail": (
                    "journal-facing DOCX contains repository path markers: "
                    + ", ".join(internal_path_markers)
                    if internal_path_markers
                    else "journal-facing DOCX contains no local repository path markers"
                ),
            },
            {
                "id": "docx:inline-figures",
                "status": "pass" if len(document.inline_shapes) >= 6 else "fail",
                "detail": f"{len(document.inline_shapes)} inline figures present",
            },
        ]
    )

    if PDF.is_file():
        pages = len(PdfReader(str(PDF)).pages)
        checks.append(
            {
                "id": "pdf:rendered-pages",
                "status": "pass" if pages == 12 else "warning",
                "detail": f"{pages} rendered PDF pages",
            }
        )
    else:
        checks.append({"id": "pdf:rendered-pages", "status": "warning", "detail": "draft PDF is absent"})

    failures = sum(item["status"] == "fail" for item in checks)
    warnings = sum(item["status"] == "warning" for item in checks)
    return {
        "schema": SCHEMA,
        "status": "fail" if failures else ("draft_verified_author_input_pending" if warnings else "draft_verified"),
        "final_upload_allowed": False,
        "checks": checks,
        "summary": {"pass": len(checks) - failures - warnings, "warning": warnings, "fail": failures},
    }


def main() -> int:
    record = audit()
    print(json.dumps(record, indent=2, ensure_ascii=False, sort_keys=True))
    return 0 if record["status"] != "fail" else 1


if __name__ == "__main__":
    raise SystemExit(main())
