"""Audit the public DataverseNO obstructed H2 release component package."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

import openpyxl


SCHEMA = "degali.dataverse-no-obstructed-release-audit.v1"
DOI = "10.18710/WSKBIJ"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def audit(root: Path) -> dict[str, Any]:
    readme = root / "00_Readme.txt"
    workbook = root / "Experimental_results_.xlsx"
    missing = [str(path.name) for path in (readme, workbook) if not path.is_file()]
    if missing:
        return {
            "schema": SCHEMA,
            "dataset_doi": DOI,
            "status": "withheld_missing_component_files",
            "promotion_allowed": False,
            "missing_files": missing,
            "concentration_validation_allowed": False,
        }
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    if "Flow parameters" not in book.sheetnames:
        raise ValueError("Flow parameters sheet is missing")
    rows = list(book["Flow parameters"].values)
    header = rows[3]
    required = (
        "Nr", "Nozzle diameter [mm]", "Pressure reservoir [bar]",
        "ignition position", "Pressure sensor_1 [kPa]", "Pressure sensor_2 [kPa]",
        "Pressure sensor_3 [kPa]", "Pressure sensor_4 [kPa]",
    )
    if tuple(header) != required:
        raise ValueError(f"unexpected workbook header: {header!r}")
    experiments = []
    experiment_row_count = 0
    for row in rows[4:]:
        if not row or row[0] in (None, "") or not str(row[0])[0].isdigit():
            continue
        experiment_row_count += 1
        numeric = row[0:8]
        if any(value in (None, "-") for value in numeric[1:]):
            continue
        experiments.append({
            "experiment": str(row[0]),
            "nozzle_diameter_mm": float(row[1]),
            "reservoir_pressure_bar": float(row[2]),
            "ignition_position": str(row[3]),
            "pressure_sensor_kpa": [float(value) for value in row[4:8]],
        })
    pressure_values = [value for item in experiments for value in item["pressure_sensor_kpa"]]
    text = readme.read_text(encoding="utf-8", errors="replace")
    has_no_ambient = "Ambient conditions (temperature, humidity, wind) were not documented." in text
    has_no_concentration = "concentration" not in text.lower() or "concentration" not in " ".join(str(value) for value in header).lower()
    return {
        "schema": SCHEMA,
        "dataset_doi": DOI,
        "status": "component_source_pressure_ready",
        "promotion_allowed": False,
        "concentration_validation_allowed": False,
        "scope": "obstructed_high_pressure_h2_source_and_explosion_pressure_component_only",
        "files": {
            "readme": {"path": str(readme), "bytes": readme.stat().st_size, "sha256": _sha256(readme)},
            "experimental_results": {"path": str(workbook), "bytes": workbook.stat().st_size, "sha256": _sha256(workbook)},
        },
        "workbook": {
            "sheet_names": book.sheetnames,
            "experiment_row_count": experiment_row_count,
            "experiment_count_with_complete_pressure_rows": len(experiments),
            "experiment_ids": [item["experiment"] for item in experiments],
            "nozzle_diameter_mm_range": [min(item["nozzle_diameter_mm"] for item in experiments), max(item["nozzle_diameter_mm"] for item in experiments)],
            "reservoir_pressure_bar_range": [min(item["reservoir_pressure_bar"] for item in experiments), max(item["reservoir_pressure_bar"] for item in experiments)],
            "pressure_sensor_kpa_range": [min(pressure_values), max(pressure_values)],
        },
        "metadata_gates": {
            "ambient_weather_missing": has_no_ambient,
            "direct_unignited_concentration_channel_present": False,
            "obstacle_geometry_machine_registered": False,
            "common_clock_declared_in_readme": "trigger signal" in text,
        },
        "gate_codes": [
            "component_source_pressure_only",
            "direct_concentration_channel_missing",
            "obstacle_geometry_not_machine_registered",
            "ambient_weather_missing" if has_no_ambient else "ambient_weather_present",
        ],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("root", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = audit(args.root)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"wrote {args.output} status={result['status']} promotion_allowed={result['promotion_allowed']}")
    return 0 if result["status"] == "component_source_pressure_ready" else 1


if __name__ == "__main__":
    raise SystemExit(main())
