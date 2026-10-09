"""Build the standalone Word Highlights file for the IJHE submission."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "IJHE_HIGHLIGHTS.md"


def read_bullets(source: Path = SOURCE) -> list[str]:
    bullets = [
        line[2:].strip()
        for line in source.read_text(encoding="utf-8").splitlines()
        if line.startswith("- ")
    ]
    if not 3 <= len(bullets) <= 5:
        raise ValueError(f"IJHE Highlights require 3-5 bullets, got {len(bullets)}")
    if any(len(item) > 85 for item in bullets):
        raise ValueError("IJHE Highlights contain a bullet longer than 85 characters")
    return bullets


def build(output: Path, source: Path = SOURCE) -> dict[str, object]:
    bullets = read_bullets(source)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    for bullet in bullets:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(8)
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.add_run(bullet)
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return {"output": str(output), "bullets": bullets, "count": len(bullets)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = build(args.output, args.source)
    print({"output": result["output"], "bullets": result["count"]})
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
