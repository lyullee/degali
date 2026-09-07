"""Frozen paired-temperature symmetry and read-only directional wind evidence.

Run using bundled Python for XLSX extraction, never modifies source workbooks.
"""
from datetime import datetime,timezone
import hashlib
import json
from pathlib import Path
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[1]
REF=ROOT/'reference/preslhy'


def symmetric_errors(left,right,prediction):
    l,r,p=np.broadcast_arrays(*[np.asarray(v,float) for v in (left,right,prediction)])
    if any(np.any(~np.isfinite(v)) for v in (l,r,p)): raise ValueError('finite inputs required')
    difference=abs(l-r)
    error=abs(p-l)+abs(p-r)
    return dict(pair_absolute_error_sum=error,pair_absolute_error_sum_floor=difference,
        pair_MAE_floor=difference/2,pair_squared_error_sum_floor=difference**2/2,
        pair_MAE_excess=(error-difference)/2)


def wind_components(speed,wind_from,release_bearing=75.):
    u,a=np.broadcast_arrays(np.asarray(speed,float),np.asarray(wind_from,float))
    if np.any(~np.isfinite(u)) or np.any(~np.isfinite(a)) or np.any(u<0) or not np.isfinite(release_bearing):
        raise ValueError('finite wind and nonnegative speed required')
    delta=np.deg2rad(a+180.-release_bearing)
    return np.stack([u*np.cos(delta),u*np.sin(delta)],axis=-1)


def seconds(value):
    t=datetime.strptime(str(value),'%H:%M:%S')
    return t.hour*3600+t.minute*60+t.second


def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()


def serial(value):
    if isinstance(value,np.ndarray): return value.tolist()
    if isinstance(value,np.generic): return value.item()
    raise TypeError(type(value).__name__)


def main():
    input_path=REF/'temperature_mean_audit_2026-09-06.json'
    data=json.loads(input_path.read_text(encoding='utf-8'))
    paths=[input_path,ROOT/'docs/prereg-lateral-temperature-symmetry.md',
        ROOT/'tests/test_lateral_temperature_symmetry.py',Path(__file__).resolve(),
        ROOT/'tmp/pdfs/PRESLHY_D3.6_Summary_of_Rainout_Experiments_V1.20.pdf']
    paths.extend(ROOT/value['path'] for value in data['workbooks'].values())
    hashes={str(p.relative_to(ROOT)):digest(p) for p in paths}
    for name,expected in data['sha256'].items():
        if digest(ROOT/name)!=expected: raise ValueError('original mean-audit input changed')
    selected=[r for r in data['rows'] if r['channel'].startswith('Left_')]
    pairs=[]; wind={}; pre={}; compass={name:22.5*i for i,name in enumerate(
        ['N','NNE','NE','ENE','E','ESE','SE','SSE','S','SSW','SW','WSW','W','WNW','NW','NNW'])}
    for left in selected:
        right=next(r for r in data['rows'] if r['trial']==left['trial'] and r['x']==left['x']
            and r['y']==-left['y'] and r['z']==left['z'] and r['channel'].startswith('Right_'))
        if left['time_seconds']!=right['time_seconds']: raise ValueError('paired timestamps differ')
        differences=np.array(left['values_K'])-right['values_K']
        row=dict(trial=left['trial'],x_m=left['x'],z_m=left['z'],y_pair_m=[left['y'],right['y']],
            channels=[left['channel'],right['channel']],ranges=[left['source_range'],right['source_range']],
            samples=len(differences),paired_difference_K=dict(mean=float(np.mean(differences)),
                median=float(np.median(differences)),p05=float(np.percentile(differences,5)),
                p95=float(np.percentile(differences,95)),std_sample=float(np.std(differences,ddof=1)),
                positive_fraction=float(np.mean(differences>0))),statistics={})
        for metric in ('minimum','p05','median','mean'):
            l,r=left['stats'][metric+'_K'],right['stats'][metric+'_K']
            row['statistics'][metric]=dict(left_K=l,right_K=r,models={})
            for model,p in left['predictions_K'].items():
                if abs(p-right['predictions_K'][model])>1e-8: raise ValueError('model is not paired-symmetric')
                row['statistics'][metric]['models'][model]=dict(prediction_K=p,**symmetric_errors(l,r,p))
        pairs.append(row)
    if len(pairs)!=4: raise ValueError('not all four frozen pairs')
    for ns,info in data['workbooks'].items():
        n=int(ns);path=ROOT/info['path'];first,last=info['rows_inclusive']
        book=openpyxl.load_workbook(path,read_only=True,data_only=True)
        sheet=book['Flexlogger'];header=list(next(sheet.iter_rows(min_row=1,max_row=1,values_only=True)))
        rows=list(sheet.iter_rows(min_row=first,max_row=last,values_only=True))
        ts=[seconds(row[header.index('time')]) for row in rows]
        if n==23:
            fast=dict(status='rejected_reported_fault_trials16_to23',source_sheet='Flexlogger',
                range=f'CT{first}:CU{last}',height_m=3.,values_not_used=True,report_page=31)
        else:
            directions=np.array([row[header.index('Wind_Direction')] for row in rows],float)
            speeds=np.array([row[header.index('Wind_Speed')] for row in rows],float)
            uv=wind_components(speeds,directions)
            vector=np.mean(uv,axis=0);unit=wind_components(np.ones_like(speeds),directions)
            mean_unit=np.mean(unit,axis=0)
            fast=dict(status='valid_report_no_fault',source_sheet='Flexlogger',range=f'CT{first}:CU{last}',
                time_range=f'A{first}:A{last}',height_m=3.,samples=len(ts),
                mean_scalar_speed_m_s=float(np.mean(speeds)),mean_velocity_source_right_m_s=vector,
                vector_mean_speed_m_s=float(np.linalg.norm(vector)),
                velocity_covariance_population=np.cov(uv,rowvar=False,ddof=0),
                mean_direction_from_deg=float((np.rad2deg(np.arctan2(vector[1],vector[0]))+75-180)%360),
                circular_mean_from_deg=float((np.rad2deg(np.arctan2(mean_unit[1],mean_unit[0]))+75-180)%360),
                circular_resultant_length=float(np.linalg.norm(mean_unit)),
                positive_right_velocity_fraction=float(np.mean(uv[:,1]>0)),
                speed_m_s=speeds,direction_from_deg=directions,time_seconds=ts,
                velocity_source_right_m_s=uv)
        local=book['LocalWeather'];all_local=list(local.iter_rows(values_only=True));lh=list(all_local[0])
        times=np.array([seconds(r[lh.index('Time')]) for r in all_local[1:]])
        within=np.flatnonzero((times>=ts[0])&(times<=ts[-1])).tolist()
        before=np.flatnonzero(times<ts[0]);after=np.flatnonzero(times>ts[-1])
        indices=sorted(set(within+([int(before[-1])] if len(before) else [])+([int(after[0])] if len(after) else [])))
        local_records=[]
        for i in indices:
            row=all_local[i+1];speed=float(row[lh.index('Loc_WindSpeedms')]);heading=row[lh.index('Loc_WindDirection')]
            angle=compass.get(heading)
            local_records.append(dict(excel_row=i+2,source_range=f'A{i+2}:K{i+2}',time=row[0],
                speed_m_s=speed,compass=heading,direction_from_deg=angle,
                velocity_source_right_m_s=None if angle is None else wind_components(speed,angle),
                missing_direction=angle is None))
        wind[ns]=dict(fast=fast,local_bracketing_records=local_records,local_height_m=1.5,
            local_period_minutes=5,local_timestamp_start_or_end_convention_confirmed=False,
            release_bearing_deg=75.,positive_right_bearing_deg=165.)
        pre_rows=list(sheet.iter_rows(min_row=2,max_row=11,values_only=True))
        pres=[]
        for pair in [p for p in pairs if p['trial']==n]:
            a,b=[header.index(c) for c in pair['channels']]
            lv=np.array([float(row[a]) for row in pre_rows]);rv=np.array([float(row[b]) for row in pre_rows])
            pres.append(dict(channels=pair['channels'],source_ranges=[f'{get_column_letter(i+1)}2:{get_column_letter(i+1)}11' for i in (a,b)],
                means_C=[float(lv.mean()),float(rv.mean())],mean_left_minus_right_K=float(np.mean(lv-rv))))
        fi=header.index('MFM1_Mass_Flow_Rate')
        pre[ns]=dict(time_range='A2:A11',first_time=pre_rows[0][0],last_time=pre_rows[-1][0],
            reported_mass_flow_g_s=[float(row[fi]) for row in pre_rows],pairs=pres,not_applied_as_correction=True)
        book.close()
    if hashes!={str(p.relative_to(ROOT)):digest(p) for p in paths}: raise ValueError('input changed')
    summary={metric:dict(paired_sensors=8,all_sensors=41,
        symmetric_41_sensor_MAE_floor_K=sum(abs(p['statistics'][metric]['left_K']-p['statistics'][metric]['right_K']) for p in pairs)/41)
        for metric in ('minimum','p05','median','mean')}
    with (REF/'lateral_temperature_symmetry_2026-09-06.json').open('x',encoding='utf-8') as stream:
        json.dump(dict(completed=True,finished_utc=datetime.now(timezone.utc).isoformat(),sha256=hashes,
            pairs=pairs,summary=summary,wind=wind,prewindow=pre,new_field=False,candidate_promoted=False,
            no_fitted_yaw=True),stream,indent=2,default=serial,allow_nan=False)
    print(json.dumps(dict(summary=summary,wind={n:{'fast':{k:v for k,v in x['fast'].items()
        if k not in ['speed_m_s','direction_from_deg','time_seconds','velocity_source_right_m_s']},
        'local':x['local_bracketing_records']} for n,x in wind.items()},prewindow=pre),indent=2,default=serial))


if __name__=='__main__': main()
