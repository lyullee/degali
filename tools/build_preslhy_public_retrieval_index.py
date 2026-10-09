"""Build a byte-level retrieval audit for public PRESLHY workbooks.

The output is reviewer-facing metadata only.  The downloaded third-party
workbooks are deliberately kept outside the repository and are never copied
into the IJHE upload bundle.  Each record requires both a public download
copy and the local workbook used by the conditional replay so that a claimed
match is an actual SHA-256 comparison rather than a filename assertion.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCHEMA = "degali.preslhy-public-file-retrieval-index.v1"
TRIAL_RE = re.compile(r"trial[_-](\d+)[_-]")
NOMINAL_SAMPLING = {
    "Flexlogger": "1 s",
    "Draeger": "1 s",
    "Xensor": "approximately 0.3 s",
    "LocalWeather": "5 min",
    "Flowmeter": "1 s",
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _relative(path: Path) -> str:
    try:
        return path.resolve().relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return path.name


def _trial_from_path(path: Path) -> int:
    match = TRIAL_RE.search(path.name.lower())
    if not match:
        raise ValueError(f"cannot infer trial number from {path.name!r}")
    return int(match.group(1))


def _mapping(values: list[str], label: str) -> dict[int, str]:
    result: dict[int, str] = {}
    for value in values:
        trial_text, separator, payload = value.partition("=")
        if not separator or not trial_text.isdigit() or not payload:
            raise ValueError(f"{label} must use TRIAL=VALUE: {value!r}")
        trial = int(trial_text)
        if trial in result:
            raise ValueError(f"duplicate {label} for trial {trial}")
        result[trial] = payload
    return result


def _sheet_records(path: Path) -> list[dict[str, object]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        records: list[dict[str, object]] = []
        for name in workbook.sheetnames:
            sheet = workbook[name]
            records.append(
                {
                    "name": name,
                    "rows": int(sheet.max_row or 0),
                    "columns": int(sheet.max_column or 0),
                    "nominal_sampling": NOMINAL_SAMPLING.get(name, "not declared"),
                }
            )
        return records
    finally:
        workbook.close()


def _record(
    trial: int,
    public_url: str,
    public_path: Path,
    local_path: Path,
) -> dict[str, object]:
    for path in (public_path, local_path):
        if not path.is_file():
            raise FileNotFoundError(path)
    public_sha = _sha256(public_path)
    local_sha = _sha256(local_path)
    public_bytes = public_path.stat().st_size
    local_bytes = local_path.stat().st_size
    return {
        "trial": trial,
        "file": {
            "name": local_path.name,
            "public_url": public_url,
            "public_bytes": public_bytes,
            "public_sha256": public_sha,
            "local_match": _relative(local_path),
            "local_bytes": local_bytes,
            "local_sha256": local_sha,
            "local_match_verified": public_bytes == local_bytes and public_sha == local_sha,
        },
        "sheets": _sheet_records(local_path),
        "promotion_boundary": {
            "calibration_certificate_present": False,
            "absolute_utc_clock_present": False,
            "surveyed_obstacle_geometry_present": False,
            "event_level_geometry_manifest_present": False,
            "headline_validation_score_changed": False,
            "evidence_status": "conditional",
            "promotion_allowed": False,
        },
    }


def build(
    output: Path,
    public_urls: dict[int, str],
    public_files: dict[int, Path],
    local_files: dict[int, Path],
) -> dict[str, object]:
    trials = sorted(set(public_urls) | set(public_files) | set(local_files))
    if not trials:
        raise ValueError("at least one retrieval record is required")
    if set(public_urls) != set(public_files) or set(public_urls) != set(local_files):
        raise ValueError("public URLs, public files and local files must cover the same trials")
    records = [
        _record(trial, public_urls[trial], public_files[trial], local_files[trial])
        for trial in trials
    ]
    payload = {
        "schema": SCHEMA,
        "retrieved_utc": datetime.now(timezone.utc).isoformat(),
        "source_record": {
            "title": "Fuel Cells and Hydrogen Joint Undertaking (FCH JU); Summary of experiment series E3.5",
            "kitopen_doi": "10.5445/IR/1000136281",
            "radar_version_doi": "10.35097/1481",
            "record_url": "https://publikationen.bibliothek.kit.edu/1000136281",
            "license": "CC BY-SA 4.0",
            "coordinate_report": {
                "name": "PRESLHY_D3.6_Summary_of_Rainout_Experiments_V1.22.pdf",
                "public_url": "https://www.radar-service.eu/radar-backend/archives/nWczysTWjmuzgFKm/retrieveFile/VCPOsvkcVmDfVaiw",
                "public_bytes": 6518995,
                "public_sha256": "865b2b9f966e05023fbb581a4ed68f95f2353d6d6db2ad137aae7a9cf17c4f67",
                "evidence_pages": {
                    "common_logging_and_sampling": 14,
                    "nominal_sensor_accuracy": 15,
                    "near_field_receptor_coordinates": [43, 44],
                    "far_field_receptor_heights": 45,
                    "weather_station_locations": 46,
                    "obstruction_offsets": 50,
                },
                "redistributed": False,
            },
        },
        "records": records,
        "promotion_rule": "Byte-level public-file matches improve traceability only; conditional bundles remain non-promotable until calibration, clock and geometry gates are closed.",
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return payload


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--public-url", action="append", required=True, help="TRIAL=PUBLIC_URL")
    parser.add_argument("--public-file", action="append", required=True, help="TRIAL=DOWNLOADED_FILE")
    parser.add_argument("--local-file", action="append", required=True, help="TRIAL=LOCAL_FILE")
    args = parser.parse_args(argv)
    payload = build(
        args.output,
        _mapping(args.public_url, "--public-url"),
        {trial: Path(value).resolve() for trial, value in _mapping(args.public_file, "--public-file").items()},
        {trial: Path(value).resolve() for trial, value in _mapping(args.local_file, "--local-file").items()},
    )
    print(json.dumps({"output": str(args.output), "records": len(payload["records"])}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
