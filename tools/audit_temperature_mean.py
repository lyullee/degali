"""Read-only same-window XLSX temperature statistics; run with bundled Python."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / 'reference/preslhy'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def statistics(values, seconds, prediction):
    a = np.asarray(values, dtype=float)
    t = np.asarray(seconds, dtype=float)
    if a.ndim != 1 or a.size < 2 or a.shape != t.shape:
        raise ValueError('matching nonempty sample vectors required')
    if not np.all(np.isfinite(a)) or not np.all(np.isfinite(t)):
        raise ValueError('nonfinite values are not silently dropped')
    if not np.allclose(np.diff(t), 1., rtol=0, atol=1e-8):
        raise ValueError('frozen window is not uniformly sampled at 1 Hz')
    mean = float(np.mean(a))
    variance = float(np.mean((a - mean)**2))
    median = float(np.median(a))
    return dict(samples=int(a.size), minimum_K=float(a.min()),
        p05_K=float(np.percentile(a, 5)), median_K=median,
        mean_K=mean, p95_K=float(np.percentile(a, 95)),
        standard_deviation_sample_K=float(np.std(a, ddof=1)),
        moment_skewness=float(np.mean((a-mean)**3) / variance**1.5) if variance else 0.,
        mean_minus_median_K=mean-median,
        trapezoid_time_mean_K=float(np.trapezoid(a, t)/(t[-1]-t[0])),
        fraction_at_or_below_CONTROL=float(np.mean(a <= prediction)),
        duration_between_endpoints_s=float(t[-1]-t[0]))


def seconds(value):
    if not isinstance(value, str):
        raise ValueError('unrecognised raw timestamp type')
    clock = datetime.strptime(value, '%H:%M:%S')
    return clock.hour*3600 + clock.minute*60 + clock.second


def main():
    paths = [REF/'stage1_single_image_2026-09-06.json',
        REF/'temperature_source_rate_audit_2026-09-05.json',
        ROOT/'docs/prereg-temperature-mean-audit.md', Path(__file__).resolve(),
        ROOT/'tests/test_temperature_mean_audit.py']
    frozen, audit = (json.loads(p.read_text(encoding='utf-8')) for p in paths[:2])
    raw_paths = {int(n): REF/'raw'/info['file'] for n, info in audit['raw_workbooks'].items()}
    paths.extend(raw_paths.values())
    hashes = {str(p.relative_to(ROOT)): digest(p) for p in paths}
    result = []; workbooks = {}
    for n, path in raw_paths.items():
        info = audit['raw_workbooks'][str(n)]
        if digest(path).lower() != info['sha256'].lower():
            raise ValueError('raw workbook differs from earlier audit')
        first, last = info['flexlogger_rows']
        with path.open('rb') as stream:
            book = openpyxl.load_workbook(stream, read_only=True, data_only=True)
            sheet = book['Flexlogger']
            header = list(next(sheet.iter_rows(min_row=1,max_row=1,values_only=True)))
            data = list(sheet.iter_rows(min_row=first,max_row=last,values_only=True))
            time_col = header.index('time')
            times = [seconds(row[time_col]) for row in data]
            rows = [r for r in frozen['temperature_rows'] if r['trial']==n]
            for old in rows:
                column = header.index(old['channel'])
                values = [float(row[column])+273.15 for row in data]
                stats = statistics(values,times,old['CONTROL_K'])
                if stats['samples'] != old['samples']:
                    raise ValueError('frozen sample count changed')
                replay = {k:abs(stats[k+'_K']-old['observed_'+k+'_K'])
                          for k in ('minimum','p05','median')}
                if max(replay.values())>1e-8:
                    raise ValueError('frozen statistics not reproduced')
                letter = get_column_letter(column+1)
                result.append(dict(trial=n,channel=old['channel'],x=old['x'],y=old['y'],z=old['z'],
                    source_workbook=str(path.relative_to(ROOT)),sheet='Flexlogger',
                    source_range=f'{letter}{first}:{letter}{last}',
                    stats=stats,original_statistics_replay_K=replay,
                    predictions_K={key:old[key] for key in ('BASE_SINGLE_IMAGE_K','CONTROL_K')},
                    values_K=values,time_seconds=times))
            workbooks[str(n)] = dict(path=str(path.relative_to(ROOT)),sha256=digest(path),
                rows_inclusive=[first,last],time_range=f'A{first}:A{last}',
                first_time=data[0][time_col],last_time=data[-1][time_col],
                interval_s=sorted(set(np.diff(times).tolist())),read_only=True)
            book.close()
    if len(result)!=41 or len({(r['trial'],r['channel']) for r in result})!=41:
        raise ValueError('not exactly original 41 unique sensors')
    summary = {}
    for subset in ('all',10,23):
        rows = result if subset=='all' else [r for r in result if r['trial']==subset]
        out = {}
        for model in ('BASE_SINGLE_IMAGE_K','CONTROL_K'):
            out[model] = {}
            for metric in ('minimum','p05','median','mean'):
                errors = np.array([r['predictions_K'][model]-r['stats'][metric+'_K'] for r in rows])
                out[model][metric] = dict(count=len(rows),MAE_K=float(np.mean(abs(errors))),
                    RMSE_K=float(np.sqrt(np.mean(errors**2))),mean_signed_error_K=float(np.mean(errors)),
                    maximum_abs_error_K=float(max(abs(errors))))
        out['mean_median_difference'] = dict(mean_abs_K=float(np.mean([abs(r['stats']['mean_minus_median_K']) for r in rows])),
            maximum_abs_K=float(max(abs(r['stats']['mean_minus_median_K']) for r in rows)))
        summary[str(subset)] = out
    if hashes!={str(p.relative_to(ROOT)):digest(p) for p in paths}:
        raise ValueError('an input changed while reading')
    output = REF/'temperature_mean_audit_2026-09-06.json'
    with output.open('x',encoding='utf-8') as stream:
        json.dump(dict(completed=True,finished_utc=datetime.now(timezone.utc).isoformat(),
            sha256=hashes,workbooks=workbooks,rows=result,summary=summary,
            frozen_statistics_unchanged=True,new_dispersion_field=False,candidate_promoted=False,
            uncertainty_interval_evaluated=False),stream,indent=2,allow_nan=False)
    print(json.dumps(summary,indent=2))


if __name__=='__main__':
    main()
