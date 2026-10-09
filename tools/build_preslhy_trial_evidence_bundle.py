"""Build a derived, hash-pinned evidence bundle for one PRESLHY E3.5 trial.

The public PRESLHY workbook is intentionally not copied into the repository.
This tool extracts only the fields needed to reconcile the source, weather,
receptor and clock channels and records the raw-input hashes.  The resulting
bundle is a *conditional qualification input*: it is never validation
promotion and it does not replace the published workbook or report.
"""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, time
import hashlib
import json
from pathlib import Path
import re
from typing import Iterable

import openpyxl

from degali.validation.preslhy import (
    _SERIAL,
    _clock_seconds,
    read_trial,
    sensor_positions,
    synchronized_gas_window,
)


CARDINAL_DEGREES = {
    "N": 0.0,
    "NNE": 22.5,
    "NE": 45.0,
    "ENE": 67.5,
    "E": 90.0,
    "ESE": 112.5,
    "SE": 135.0,
    "SSE": 157.5,
    "S": 180.0,
    "SSW": 202.5,
    "SW": 225.0,
    "WSW": 247.5,
    "W": 270.0,
    "WNW": 292.5,
    "NW": 315.0,
    "NNW": 337.5,
}


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _clock_value(value: object) -> float:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.hour * 3600.0 + value.minute * 60.0 + value.second + value.microsecond / 1.0e6
    if isinstance(value, str):
        match = re.fullmatch(r"\s*(\d{1,2}):(\d{2}):(\d{2}(?:\.\d+)?)\s*", value)
        if match:
            hour, minute, second = match.groups()
            return 3600.0 * float(hour) + 60.0 * float(minute) + float(second)
    raise ValueError(f"unsupported clock value: {value!r}")


def _finite(value: object) -> float | None:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if number == number and abs(number) != float("inf") else None


def _header_index(header: Iterable[object], contains: str) -> int:
    for index, value in enumerate(header):
        if value and contains.lower() in str(value).lower():
            return index
    raise KeyError(f"missing heading containing {contains!r}")


def _write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, object]]) -> int:
    count = 0
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)
            count += 1
    return count


def build_bundle(workbook_path: Path, report_path: Path, output_dir: Path, *, event_id: str, operator_id: str, sensor_set_id: str) -> dict[str, object]:
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)
    if not report_path.is_file():
        raise FileNotFoundError(report_path)
    output_dir.mkdir(parents=True, exist_ok=True)

    positions = sensor_positions(report_path)
    trial = read_trial(workbook_path, positions, gas_transport_delay_s=18.0)
    book = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)

    flow_rows = list(book["Flowmeter"].iter_rows(values_only=True))
    flow_header = flow_rows[0]
    flow_time_index = 0
    flow_rate_index = _header_index(flow_header, "MassFlow")
    flow_clock = _clock_seconds([row[flow_time_index] for row in flow_rows[1:]])
    event_clock_s = float(flow_clock[0])
    source_id = f"{event_id}-source"
    weather_id = f"{event_id}-weather"
    obstacle_id = f"{event_id}-open-pad-geometry"
    receptor_id = f"{event_id}-receptors"
    common_clock_id = f"{event_id}-shared-wall-clock"

    source_rows = []
    for row, clock_s in zip(flow_rows[1:], flow_clock):
        raw_rate = _finite(row[flow_rate_index])
        if raw_rate is None:
            continue
        source_rows.append({
            "event_id": event_id,
            "source_boundary_id": source_id,
            "common_clock_id": common_clock_id,
            "timestamp": str(row[flow_time_index]),
            "time_s": f"{clock_s - event_clock_s:.6f}",
            "mass_flow_g_s_raw": f"{raw_rate:.9g}",
            "mass_flow_kg_s": f"{max(raw_rate, 0.0) / 1000.0:.9g}",
            "phase": "LH2",
            "release_mode": "horizontal elevated release",
            "release_height_m": "0.5",
            "nozzle_diameter_mm": "25.4",
            "source_coordinate_reference": "E3.5 source-relative facility coordinates",
        })
    source_count = _write_csv(
        output_dir / "source_boundary.csv",
        list(source_rows[0].keys()), source_rows,
    )

    weather_sheet = list(book["LocalWeather"].iter_rows(values_only=True))
    weather_header = weather_sheet[0]
    weather_time_index = _header_index(weather_header, "Time")
    weather_speed_index = _header_index(weather_header, "WindSpeed")
    weather_direction_index = _header_index(weather_header, "WindDirection")
    weather_temperature_index = _header_index(weather_header, "OutdoorTemperature")
    weather_humidity_index = _header_index(weather_header, "OutdoorHumidity")
    weather_pressure_index = _header_index(weather_header, "RelativePressure")
    weather_dew_point_index = _header_index(weather_header, "DewPoint")
    weather_rows = []
    weather_clock = _clock_seconds([row[weather_time_index] for row in weather_sheet[1:]])
    for row, clock_s in zip(weather_sheet[1:], weather_clock):
        speed = _finite(row[weather_speed_index])
        if speed is None:
            continue
        direction_text = str(row[weather_direction_index]).strip().upper()
        direction_deg = CARDINAL_DEGREES.get(direction_text)
        weather_rows.append({
            "event_id": event_id,
            "weather_id": weather_id,
            "common_clock_id": common_clock_id,
            "timestamp": str(row[weather_time_index]),
            "time_s": f"{clock_s - event_clock_s:.6f}",
            "wind_speed_m_s": f"{speed:.9g}",
            "wind_direction": direction_text,
            "wind_direction_deg": "" if direction_deg is None else f"{direction_deg:.9g}",
            "measurement_interval_s": "300",
            "measurement_height_m": "",
            "outdoor_temperature_c": "" if _finite(row[weather_temperature_index]) is None else f"{_finite(row[weather_temperature_index]):.9g}",
            "outdoor_humidity_pct": "" if _finite(row[weather_humidity_index]) is None else f"{_finite(row[weather_humidity_index]):.9g}",
            "relative_pressure_hpa": "" if _finite(row[weather_pressure_index]) is None else f"{_finite(row[weather_pressure_index]):.9g}",
            "dew_point_c": "" if _finite(row[weather_dew_point_index]) is None else f"{_finite(row[weather_dew_point_index]):.9g}",
            "stability_status": "not_provided_in_public_workbook",
        })
    weather_count = _write_csv(
        output_dir / "weather.csv",
        list(weather_rows[0].keys()), weather_rows,
    )

    gas_rows = list(book["Xensor"].iter_rows(values_only=True))
    gas_header = gas_rows[0]
    gas_clock = _clock_seconds([row[0] for row in gas_rows[1:]])
    aligned = synchronized_gas_window(
        flow_clock, gas_clock, trial.window, gas_transport_delay_s=18.0,
    )
    if aligned is None:
        raise RuntimeError("PRESLHY source and Xensor clocks could not be aligned")
    gas_lo, gas_hi = aligned
    receptor_rows = []
    mapped_sensor_count = 0
    for column, heading in enumerate(gas_header[1:], start=1):
        if not heading:
            continue
        match = _SERIAL.search(str(heading))
        if match is None or match.group(1) not in positions:
            continue
        serial = match.group(1)
        x_m, y_m, z_m = positions[serial]
        mapped_sensor_count += 1
        for row, clock_s in zip(gas_rows[1:][gas_lo:gas_hi], gas_clock[gas_lo:gas_hi]):
            raw = _finite(row[column] if column < len(row) else None)
            if raw is None:
                continue
            receptor_rows.append({
                "event_id": event_id,
                "receptor_geometry_id": receptor_id,
                "sensor_id": serial,
                "sensor_set_id": sensor_set_id,
                "common_clock_id": common_clock_id,
                "timestamp": str(row[0]),
                "time_s": f"{clock_s - event_clock_s:.6f}",
                "x_m": f"{x_m:.9g}",
                "y_m": f"{y_m:.9g}",
                "height_m": f"{z_m:.9g}",
                "observed_vol_pct_raw": f"{raw:.9g}",
                "observed_vol_pct": f"{max(raw, 0.0):.9g}",
                "sensor_processing": "negative baseline values retained in raw column; physical concentration clipped at zero",
            })
    receptor_count = _write_csv(
        output_dir / "receptor_observations.csv",
        list(receptor_rows[0].keys()), receptor_rows,
    )

    registry_rows = [
        {
            "sensor_set_id": sensor_set_id,
            "sensor_id": serial,
            "x_m": f"{x:.9g}",
            "y_m": f"{y:.9g}",
            "height_m": f"{z:.9g}",
            "response_interval_s": "0.3",
            "calibration_status": "not_provided_in_public_workbook",
            "sensor_type": "NREL Xensor 5320",
        }
        for serial, (x, y, z) in sorted(positions.items())
        if any(row["sensor_id"] == serial for row in receptor_rows)
    ]
    registry_count = _write_csv(
        output_dir / "sensor_registry.csv",
        list(registry_rows[0].keys()), registry_rows,
    )

    source_start_s = float(flow_clock[trial.window[0]] - event_clock_s)
    source_end_s = float(flow_clock[trial.window[1] - 1] - event_clock_s)
    clock_record = {
        "schema": "degali.preslhy-common-clock-record.v1",
        "event_id": event_id,
        "common_clock_id": common_clock_id,
        "synchronization_method": "shared wall-clock timestamps in Flowmeter, Xensor and LocalWeather workbook sheets",
        "timezone": "local facility time; UTC offset not recorded in public workbook",
        "clock_drift": "not provided in public workbook",
        "event_clock_origin": str(flow_rows[1][flow_time_index]),
        "source_window": {"start_s": source_start_s, "end_s": source_end_s, "samples": trial.window[1] - trial.window[0]},
        "xensor_window": {"start_index": gas_lo, "end_index": gas_hi, "transport_delay_s": 18.0},
        "weather_sampling_interval_s": 300.0,
        "operator_note": "The 18 s Xensor transport delay is declared from the PRESLHY report; it is not fitted to concentration peaks.",
    }
    (output_dir / "common_clock.json").write_text(
        json.dumps(clock_record, indent=2, ensure_ascii=False) + "\n", encoding="utf-8",
    )

    obstacle_record = {
        "schema": "degali.obstacle-geometry-record.v1",
        "event_id": event_id,
        "obstacle_geometry_id": obstacle_id,
        "coordinate_reference": "source-relative metres; z above concrete pad",
        "obstacles_present": False,
        "geometry": [],
        "basis": "E3.5 Trial 10 is represented as an open concrete-pad release; no engineered obstacle/wake geometry is promoted from the public workbook.",
        "promotion_status": "conditional_open_pad_only",
    }
    (output_dir / "obstacle_geometry.json").write_text(
        json.dumps(obstacle_record, indent=2, ensure_ascii=False) + "\n", encoding="utf-8",
    )

    raw_inputs = [
        {"role": "public_workbook", "path": str(workbook_path), "sha256": _sha256(workbook_path)},
        {"role": "sensor_coordinate_report", "path": str(report_path), "sha256": _sha256(report_path)},
    ]
    metadata = {
        "schema": "degali.preslhy-derived-evidence-bundle.v1",
        "event_id": event_id,
        "trial": trial.number,
        "source_boundary_id": source_id,
        "weather_id": weather_id,
        "obstacle_geometry_id": obstacle_id,
        "receptor_geometry_id": receptor_id,
        "common_clock_id": common_clock_id,
        "sensor_set_id": sensor_set_id,
        "operator_id": operator_id,
        "trial_window_rows": list(trial.window),
        "flow_mean_g_s": trial.flow_mean,
        "flow_peak_g_s": trial.flow_peak,
        "mapped_sensor_count": mapped_sensor_count,
        "source_row_count": source_count,
        "weather_row_count": weather_count,
        "receptor_row_count": receptor_count,
        "sensor_registry_row_count": registry_count,
        "raw_inputs": raw_inputs,
        "public_dataset_doi": "10.35097/1481",
        "promotion_allowed": False,
        "status": "conditional",
        "limitations": [
            "public workbook does not provide channel-level calibration certificate in this bundle",
            "LocalWeather is a five-minute atmospheric record, not a high-frequency mast trace",
            "obstacle geometry is an explicit open-pad declaration, not a measured 3-D site survey",
            "negative Xensor baseline values are preserved separately and clipped only in the physical concentration field",
        ],
    }
    (output_dir / "bundle_metadata.json").write_text(
        json.dumps(metadata, indent=2, ensure_ascii=False) + "\n", encoding="utf-8",
    )
    return metadata


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--event-id", default="preslhy-e35-trial-10-2019-09-13")
    parser.add_argument("--operator-id", default="degali-reproducible-bundle-builder-v1")
    parser.add_argument("--sensor-set-id", default="preslhy-e35-xensor-trial10-a1")
    args = parser.parse_args(argv)
    metadata = build_bundle(
        args.workbook.resolve(), args.report.resolve(), args.output_dir.resolve(),
        event_id=args.event_id, operator_id=args.operator_id,
        sensor_set_id=args.sensor_set_id,
    )
    print(json.dumps({"output_dir": str(args.output_dir), "status": metadata["status"], "promotion_allowed": metadata["promotion_allowed"], "mapped_sensor_count": metadata["mapped_sensor_count"]}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
