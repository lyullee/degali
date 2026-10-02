"""PRESLHY E3.4 pool mass and substrate-temperature observation operator.

E3.4 workbooks are optional local validation inputs (DOI 10.35097/1319).  No
measurement is shipped with DEGALI.  The reader extracts the experimenter's
corrected ``m(LH2) [g]`` channel rather than reconstructing pool inventory
from the total facility scale.  This matters for sand/gravel records where
solid displacement or material loss is visible in the total-weight signal.

The code intentionally requires a caller-declared evaporation interval.  The
record includes filling, sloshing and, on some substrates, mass artefacts;
selecting a favourable interval automatically from the mass slope would turn
the validation operator into a fitted source term.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import math

import numpy as np


@dataclass(frozen=True)
class E34PoolHistory:
    """Corrected pool mass plus co-recorded thermocouple histories."""

    path: Path
    time_s: np.ndarray
    liquid_mass_kg: np.ndarray
    thermocouple_k: dict[str, np.ndarray]
    substrate: str


@dataclass(frozen=True)
class E34EvaporationWindow:
    """Mass-loss observation over a fixed, caller-declared stable interval."""

    start_s: float
    end_s: float
    samples: int
    initial_mass_kg: float
    final_mass_kg: float
    evaporation_rate_kg_s: float
    mass_slope_r_squared: float
    temperature_span_k: float
    notes: tuple[str, ...]

    @property
    def mass_lost_kg(self) -> float:
        return self.initial_mass_kg - self.final_mass_kg


def _finite(values) -> np.ndarray:
    out = []
    for value in values:
        try:
            value = float(value)
        except (TypeError, ValueError):
            value = float("nan")
        out.append(value if math.isfinite(value) else float("nan"))
    return np.asarray(out, dtype=float)


def read_e34_pool_history(workbook: str | Path) -> E34PoolHistory:
    """Read one raw E3.4 workbook without altering its time origin.

    The first sheet contains the thermocouples and the scale-derived pool
    record.  The required headings live on row seven in the public files;
    locating them by name is safer than hard-coding their columns.  Mass is
    converted from the supplied corrected grams to kilograms.
    """

    import openpyxl

    path = Path(workbook)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        header = next(sheet.iter_rows(min_row=7, max_row=7, values_only=True))
        try:
            time_column = next(
                i for i, value in enumerate(header)
                if str(value).strip() == "Sync. Time [s]"
            )
        except StopIteration:
            # Concrete01 is the lone public file without a synchronised
            # export column; its thermocouple sample index is exactly seconds.
            time_column = 0
        try:
            mass_column = next(
                i for i, value in enumerate(header)
                if str(value).strip() == "m(LH2) [g]"
            )
        except StopIteration as error:
            raise ValueError(f"{path.name}: corrected m(LH2) channel missing") from error

        tc_columns = {
            str(value).strip(): i
            for i, value in enumerate(header)
            if value is not None and i < 29 and str(value).strip() not in {"", "X_Value"}
        }
        rows = list(sheet.iter_rows(min_row=8, values_only=True))
    finally:
        book.close()

    time_s = _finite([row[time_column] if time_column < len(row) else None for row in rows])
    mass_kg = 1.0e-3 * _finite([row[mass_column] if mass_column < len(row) else None for row in rows])
    valid = np.isfinite(time_s) & np.isfinite(mass_kg)
    if np.count_nonzero(valid) < 3:
        raise ValueError(f"{path.name}: fewer than three synchronized pool-mass samples")
    time_s, mass_kg = time_s[valid], mass_kg[valid]
    if np.any(np.diff(time_s) <= 0.0):
        raise ValueError(f"{path.name}: pool time is not strictly increasing")
    thermocouple_k = {
        name: _finite([row[column] if column < len(row) else None for row in rows])[valid]
        for name, column in tc_columns.items()
    }
    return E34PoolHistory(
        path=path, time_s=time_s, liquid_mass_kg=mass_kg,
        thermocouple_k=thermocouple_k,
        substrate=path.name.split("-", 2)[1].lower() if "-" in path.name else "unknown",
    )


def evaporation_window(
    history: E34PoolHistory,
    *,
    start_s: float,
    end_s: float,
    thermocouples: tuple[str, ...] = (),
) -> E34EvaporationWindow:
    """Reduce a predeclared post-fill mass-loss interval.

    A negative fitted slope is required.  ``temperature_span_k`` records the
    maximum span of selected thermal channels; it is a screening diagnostic,
    not an automatic acceptance threshold because each substrate and depth
    has a physically different cooling transient.
    """

    if end_s <= start_s:
        raise ValueError("end_s must be later than start_s")
    selected = (history.time_s >= start_s) & (history.time_s <= end_s)
    if np.count_nonzero(selected) < 3:
        raise ValueError("evaporation window needs at least three samples")
    time = history.time_s[selected]
    mass = history.liquid_mass_kg[selected]
    slope, intercept = np.polyfit(time, mass, 1)
    predicted = slope * time + intercept
    variance = float(np.sum((mass - mass.mean()) ** 2))
    r2 = 1.0 - float(np.sum((mass - predicted) ** 2)) / max(variance, 1.0e-30)
    if slope >= 0.0:
        raise ValueError("selected interval has no net evaporative mass loss")
    requested = thermocouples or tuple(history.thermocouple_k)
    missing = [name for name in requested if name not in history.thermocouple_k]
    if missing:
        raise KeyError(f"thermocouples not present: {missing}")
    spans = [
        float(np.nanmax(history.thermocouple_k[name][selected]) - np.nanmin(history.thermocouple_k[name][selected]))
        for name in requested
        if np.any(np.isfinite(history.thermocouple_k[name][selected]))
    ]
    notes = [
        "rate is the least-squares slope of the supplied corrected m(LH2) channel",
        "the time interval is caller-declared; no interval was selected to improve agreement",
    ]
    if history.substrate in {"sand", "gravel"}:
        notes.append("sand/gravel mass artefacts remain a manual review item")
    return E34EvaporationWindow(
        start_s=float(time[0]), end_s=float(time[-1]), samples=int(time.size),
        initial_mass_kg=float(mass[0]), final_mass_kg=float(mass[-1]),
        evaporation_rate_kg_s=float(-slope), mass_slope_r_squared=float(r2),
        temperature_span_k=max(spans, default=float("nan")), notes=tuple(notes),
    )


__all__ = [
    "E34PoolHistory", "E34EvaporationWindow", "read_e34_pool_history", "evaporation_window",
]
