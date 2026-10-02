"""Read the public PRESLHY E3.1 high-pressure hydrogen workbooks safely.

PRESLHY E3.1 part A is a public 80 K / 300 K *compressed-hydrogen* jet
campaign, not a liquid-hydrogen spill experiment.  This small reader exposes
only the two signals that the campaign documentation supports as a first
thermal-validation input:

* vessel/nozzle pressure and the common valve-relay signal; and
* the externally mounted axial thermocouples T5, T6 and T7.

It deliberately does not reduce the hydrogen concentration channels.  They
were sampled through 2.55 m tubing and the report documents a nominal, but
pressure-dependent, delay.  Likewise it never applies the approximately 7 K
cold-bath offset observed for the closed thermocouples as a universal
calibration.  Callers must carry that sensitivity explicitly.

The external workbooks are not package data.  ``openpyxl`` is consequently a
test/validation extra rather than a runtime dependency of :mod:`degali`.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np


_FAST_XML_MINIMUM_BYTES = 16_000_000


@dataclass(frozen=True)
class E31PressureHistory:
    """The synchronized source-pressure trace, in bar and seconds."""

    time_s: np.ndarray
    vessel_pressure_bar: np.ndarray
    nozzle_pressure_bar: np.ndarray
    valve_relay_v: np.ndarray


@dataclass(frozen=True)
class E31AxisTemperatureHistory:
    """Uncorrected closed-thermocouple records at 0.25, 0.75 and 1.75 m."""

    time_s: np.ndarray
    t5_k: np.ndarray
    t6_k: np.ndarray
    t7_k: np.ndarray
    axial_distances_m: tuple[float, float, float] = (0.25, 0.75, 1.75)
    cold_bath_offset_applied_k: float = 0.0


@dataclass(frozen=True)
class E31ReleaseLineTemperatureHistory:
    """Optional welded T4 record from the release line, in recorded Kelvin."""

    time_s: np.ndarray
    temperature_k: np.ndarray
    workbook_heading: str


@dataclass(frozen=True)
class E31HighPressureRun:
    """One public E3.1 run, without concentration or inferred particle data."""

    path: Path
    nozzle_diameter_mm: float
    initial_vessel_pressure_bar: float
    pressure: E31PressureHistory
    axis_temperature: E31AxisTemperatureHistory
    release_line_temperature: E31ReleaseLineTemperatureHistory | None


@dataclass(frozen=True)
class E31PressureRun:
    """Pressure-only E3.1 view for fast source-term validation iterations."""

    path: Path
    nozzle_diameter_mm: float
    initial_vessel_pressure_bar: float
    pressure: E31PressureHistory


@dataclass(frozen=True)
class E31ValveInterval:
    """Longest relay-open interval found from the recorded valve signal."""

    start_s: float
    end_s: float
    threshold_v: float
    samples: int


@dataclass(frozen=True)
class E31NozzlePressureRise:
    """First declared PNoz rise after the relay signal.

    The E3.1 report cautions that relay-to-nozzle delay is not constant.  This
    object records a *measurement synchronization* result, not a valve-flow
    model: the caller declares the pressure rise that constitutes response.
    """

    relay_interval: E31ValveInterval
    baseline_pressure_bar: float
    required_rise_bar: float
    threshold_pressure_bar: float
    response_time_s: float
    delay_after_relay_s: float


@dataclass(frozen=True)
class E31AxisTemperatureEnvelope:
    """Uncorrected T5/T6/T7 ranges over one explicitly declared interval."""

    interval: E31ValveInterval
    samples: int
    minimum_k: tuple[float, float, float]
    percentile_05_k: tuple[float, float, float]
    median_k: tuple[float, float, float]


@dataclass(frozen=True)
class E31IdealChokedSourceBound:
    """A pure-H2, ``Cd=1`` upper source bound derived from one E3.1 run.

    It is useful for rejecting an impossible source/thermal comparison. It is
    not a measured mass-flow boundary and is never selected by the default
    LH2 model.
    """

    storage_temperature_k: float
    storage_pressure_bar: float
    orifice_diameter_mm: float
    ideal_mass_flow_kg_s: float
    throat_pressure_bar: float
    throat_temperature_k: float
    throat_velocity_m_s: float
    atmospheric_diameter_m: float
    atmospheric_temperature_k: float
    atmospheric_density_kg_m3: float
    atmospheric_velocity_m_s: float


@dataclass(frozen=True)
class E31AxisTemperatureComparison:
    """A declared model-temperature projection against a raw E3.1 envelope.

    The experiment's thermocouple trace is transient while the integral jet
    supplies a stationary centreline.  These differences are consequently
    descriptive residuals, not a fitted objective or an accuracy score.
    """

    distances_m: tuple[float, float, float]
    model_temperature_k: tuple[float, float, float]
    observed_minimum_k: tuple[float, float, float]
    observed_percentile_05_k: tuple[float, float, float]
    observed_median_k: tuple[float, float, float]
    residual_to_percentile_05_k: tuple[float, float, float]
    inside_recorded_minimum_median: tuple[bool, bool, bool]


@dataclass(frozen=True)
class E31BlowdownPressureComparison:
    """Pressure/temperature residuals at predeclared source-time stations."""

    source_start_time_s: float
    time_after_source_start_s: tuple[float, ...]
    measured_pressure_bar: tuple[float, ...]
    model_pressure_bar: tuple[float, ...]
    pressure_residual_bar: tuple[float, ...]
    model_tank_temperature_k: tuple[float, ...]


@dataclass(frozen=True)
class E31VesselInventoryChange:
    """Constant-temperature pure-H2 inventory change over one relay interval.

    This is intentionally an integrated mass boundary. Differentiating a
    high-rate pressure record would require a separately declared smoothing
    and valve model, neither of which is implied by the raw workbook.
    """

    interval: E31ValveInterval
    vessel_volume_m3: float
    storage_temperature_k: float
    samples: int
    initial_mass_kg: float
    final_mass_kg: float
    released_mass_kg: float
    mean_release_rate_kg_s: float


@dataclass(frozen=True)
class E31InventoryReleaseTiming:
    """When a declared fraction of a pressure-inferred mass loss has occurred."""

    interval: E31ValveInterval
    storage_temperature_k: float
    vessel_volume_m3: float
    released_fraction: float
    time_after_open_s: float
    released_mass_kg: float
    total_released_mass_kg: float


def _openpyxl() -> Any:
    try:
        import openpyxl
    except ImportError as error:  # pragma: no cover - depends on installation
        raise ImportError(
            "Reading PRESLHY E3.1 workbooks needs openpyxl. "
            "Install DEGALI with its test/validation dependencies."
        ) from error
    return openpyxl


def _only_sheet(book: Any, suffix: str) -> Any:
    matches = [name for name in book.sheetnames if name.endswith(suffix)]
    if len(matches) != 1:
        raise ValueError(
            f"expected exactly one E3.1 sheet ending {suffix!r}; found {matches}"
        )
    return book[matches[0]]


def _headers(sheet: Any) -> dict[str, int]:
    values = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
    return {
        str(value).strip(): index
        for index, value in enumerate(values)
        if value not in (None, "")
    }


def _require_headers(sheet: Any, names: tuple[str, ...]) -> dict[str, int]:
    index = _headers(sheet)
    missing = [name for name in names if name not in index]
    if missing:
        raise ValueError(
            f"{sheet.title}: missing required E3.1 columns {missing}"
        )
    return index


def _release_line_heading(headings: dict[str, int]) -> str | None:
    """Find the welded T4 heading despite workbook character encodings.

    The original E3.1 exports label T4 as ``T-Stück``.  Some XLSX readers
    expose the umlaut as a replacement sequence, but the stable ``T-St``
    prefix remains.  No guessed numeric column position is used.
    """
    matches = [name for name in headings if name.casefold().startswith("t-st")]
    if len(matches) > 1:
        raise ValueError(f"ambiguous E3.1 release-line thermocouples {matches}")
    return matches[0] if matches else None


def _number(value: object) -> float:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return float("nan")
    return number if np.isfinite(number) else float("nan")


def _history(sheet: Any, headings: dict[str, int], names: tuple[str, ...]) -> dict[str, np.ndarray]:
    columns = {name: [] for name in names}
    # E3.1 sheets carry several unrelated diagnostic columns.  Streaming only
    # through the last requested channel avoids materialising those channels
    # on the 662,000-row pressure traces while retaining their original
    # zero-based header indices.
    last_column = max(headings[name] for name in names) + 1
    started = False
    consecutive_blank_time_rows = 0
    for row in sheet.iter_rows(min_row=7, max_col=last_column, values_only=True):
        time_value = _number(row[headings["Time [s]"]] if headings["Time [s]"] < len(row) else None)
        if not np.isfinite(time_value):
            if started:
                consecutive_blank_time_rows += 1
                # The public workbooks retain a large formatted Excel tail
                # after their real acquisition rows.  This is not a sampled
                # gap; ending after a long blank tail avoids scanning hundreds
                # of thousands of empty rows while tolerating short gaps.
                if consecutive_blank_time_rows >= 512:
                    break
            continue
        started = True
        consecutive_blank_time_rows = 0
        for name in names:
            column = headings[name]
            columns[name].append(time_value if name == "Time [s]" else _number(row[column] if column < len(row) else None))
    time = np.asarray(columns["Time [s]"], dtype=float)
    valid = np.isfinite(time)
    if int(np.count_nonzero(valid)) < 2:
        raise ValueError(f"{sheet.title}: fewer than two finite time samples")
    return {
        name: np.asarray(values, dtype=float)[valid]
        for name, values in columns.items()
    }


def _excel_column_name(index: int) -> str:
    """Return a zero-based Excel column label without importing openpyxl APIs."""
    if index < 0:
        raise ValueError("Excel column index must be non-negative")
    out = ""
    value = index + 1
    while value:
        value, remainder = divmod(value - 1, 26)
        out = chr(ord("A") + remainder) + out
    return out


def _fast_numeric_history(
    path: Path, sheet: Any, headings: dict[str, int], names: tuple[str, ...]
) -> dict[str, np.ndarray] | None:
    """Read a large numeric XLSX XML sheet without materializing row objects.

    The public E3.1 pressure workbooks contain a 190 MB uncompressed XML
    sheet. ``openpyxl`` is appropriate for sheet discovery and headers, but
    creates hundreds of thousands of Python row/cell objects merely to obtain
    four numeric traces. For only large sheets, stream the required XML cell
    columns using the standard library and require exact row alignment. Any
    structural deviation falls back to the conservative openpyxl reader.
    """
    import re
    import zipfile

    member = getattr(sheet, "_worksheet_path", None)
    if not isinstance(member, str):
        return None
    try:
        with zipfile.ZipFile(path) as archive:
            info = archive.getinfo(member)
            if info.file_size < _FAST_XML_MINIMUM_BYTES:
                return None
            xml = archive.read(member)
    except (KeyError, OSError, zipfile.BadZipFile):
        return None

    records: dict[str, tuple[np.ndarray, np.ndarray]] = {}
    for name in names:
        column = _excel_column_name(headings[name]).encode("ascii")
        pattern = re.compile(
            rb'<c r="' + column
            + rb'(\d+)"[^>]*>(?:(?!</c>).)*?<v>([^<]+)</v></c>',
            re.DOTALL,
        )
        rows = np.fromiter(
            (int(match.group(1)) for match in pattern.finditer(xml)),
            dtype=np.int64,
        )
        values = np.fromiter(
            (_number(match.group(2)) for match in pattern.finditer(xml)),
            dtype=float,
            count=rows.size,
        )
        data = rows >= 7
        rows, values = rows[data], values[data]
        if rows.size < 2 or values.size != rows.size:
            return None
        records[name] = (rows, values)
    time_rows, time = records["Time [s]"]
    if any(not np.array_equal(rows, time_rows) for rows, _ in records.values()):
        return None
    valid = np.isfinite(time)
    if int(np.count_nonzero(valid)) < 2:
        return None
    return {name: values[valid] for name, (_rows, values) in records.items()}


def _history_from_workbook(
    path: Path, sheet: Any, headings: dict[str, int], names: tuple[str, ...]
) -> dict[str, np.ndarray]:
    """Use the fast numeric reader when safe, otherwise preserve workbook semantics."""
    return _fast_numeric_history(path, sheet, headings, names) or _history(
        sheet, headings, names
    )


def _metadata(pressure_sheet: Any) -> tuple[float, float]:
    first = next(pressure_sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    second = next(pressure_sheet.iter_rows(min_row=2, max_row=2, values_only=True))
    labels = {
        str(value).strip(): index
        for index, value in enumerate(first)
        if value not in (None, "")
    }
    for label in ("dNoz [mm]", "Pini [bar]"):
        if label not in labels:
            raise ValueError(f"{pressure_sheet.title}: missing metadata {label!r}")
    nozzle = _number(second[labels["dNoz [mm]"]])
    initial_pressure = _number(second[labels["Pini [bar]"]])
    if not np.isfinite(nozzle) or nozzle <= 0.0:
        raise ValueError(f"{pressure_sheet.title}: invalid nozzle diameter {nozzle!r}")
    if not np.isfinite(initial_pressure) or initial_pressure <= 0.0:
        raise ValueError(
            f"{pressure_sheet.title}: invalid initial pressure {initial_pressure!r}"
        )
    return nozzle, initial_pressure


def read_e31_high_pressure_run(workbook: str | Path) -> E31HighPressureRun:
    """Read the E3.1 pressure and axial-temperature signals from one workbook.

    The returned temperatures are the recorded Kelvin values.  No correction
    is applied for the report's approximately +7 K closed-thermocouple
    liquid-nitrogen-bath observation, and the records are not suitable for
    concentration-arrival or particle-slip fitting.
    """
    path = Path(workbook)
    if not path.is_file():
        raise FileNotFoundError(path)
    openpyxl = _openpyxl()
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        pressure_sheet = _only_sheet(book, "-Press")
        temperature_sheet = _only_sheet(book, "-Temp")
        nozzle_mm, initial_bar = _metadata(pressure_sheet)
        pressure_headings = _require_headers(
            pressure_sheet,
            ("Time [s]", "Druck_1", "Druck_2", "Valve-Relay"),
        )
        temperature_headings = _require_headers(
            temperature_sheet,
            ("Time [s]", "T5", "T6", "T7"),
        )
        release_heading = _release_line_heading(_headers(temperature_sheet))
        pressure_values = _history_from_workbook(
            path, pressure_sheet, pressure_headings,
            ("Time [s]", "Druck_1", "Druck_2", "Valve-Relay"),
        )
        temperature_names = ("Time [s]", "T5", "T6", "T7")
        all_temperature_headings = dict(temperature_headings)
        if release_heading is not None:
            all_temperature_headings["T4"] = _headers(temperature_sheet)[release_heading]
            temperature_names += ("T4",)
        temperature_values = _history(
            temperature_sheet, all_temperature_headings, temperature_names
        )
    finally:
        book.close()
    return E31HighPressureRun(
        path=path,
        nozzle_diameter_mm=nozzle_mm,
        initial_vessel_pressure_bar=initial_bar,
        pressure=E31PressureHistory(
            time_s=pressure_values["Time [s]"],
            vessel_pressure_bar=pressure_values["Druck_1"],
            nozzle_pressure_bar=pressure_values["Druck_2"],
            valve_relay_v=pressure_values["Valve-Relay"],
        ),
        axis_temperature=E31AxisTemperatureHistory(
            time_s=temperature_values["Time [s]"],
            t5_k=temperature_values["T5"],
            t6_k=temperature_values["T6"],
            t7_k=temperature_values["T7"],
        ),
        release_line_temperature=(
            E31ReleaseLineTemperatureHistory(
                time_s=temperature_values["Time [s]"],
                temperature_k=temperature_values["T4"],
                workbook_heading=release_heading,
            )
            if release_heading is not None
            else None
        ),
    )


def read_e31_pressure_run(workbook: str | Path) -> E31PressureRun:
    """Read only the E3.1 ``Press`` sheet for repeatable blowdown checks.

    The full workbook reader remains required for thermal comparison.  This
    lighter view exists because the public high-rate Excel files retain large
    formatted temperature-sheet tails that do not enter a tank source-term
    validation.
    """
    path = Path(workbook)
    if not path.is_file():
        raise FileNotFoundError(path)
    openpyxl = _openpyxl()
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        pressure_sheet = _only_sheet(book, "-Press")
        nozzle_mm, initial_bar = _metadata(pressure_sheet)
        headings = _require_headers(
            pressure_sheet,
            ("Time [s]", "Druck_1", "Druck_2", "Valve-Relay"),
        )
        values = _history_from_workbook(
            path, pressure_sheet, headings,
            ("Time [s]", "Druck_1", "Druck_2", "Valve-Relay"),
        )
    finally:
        book.close()
    return E31PressureRun(
        path=path,
        nozzle_diameter_mm=nozzle_mm,
        initial_vessel_pressure_bar=initial_bar,
        pressure=E31PressureHistory(
            time_s=values["Time [s]"],
            vessel_pressure_bar=values["Druck_1"],
            nozzle_pressure_bar=values["Druck_2"],
            valve_relay_v=values["Valve-Relay"],
        ),
    )


def valve_open_interval(
    run: E31HighPressureRun | E31PressureRun,
    *,
    threshold_fraction: float = 0.5,
) -> E31ValveInterval:
    """Return the longest interval above a robust relay-amplitude threshold.

    The threshold is the requested fraction between the 5th and 95th
    percentiles of the recorded relay voltage.  This makes the criterion
    reproducible without assuming that the nominal relay high value is 5 V,
    and longest-run selection prevents an isolated electrical spike from
    defining the release.  It is a source synchronization interval, not a
    concentration-arrival time.
    """
    fraction = float(threshold_fraction)
    if not 0.0 < fraction < 1.0:
        raise ValueError("threshold_fraction must be strictly between zero and one")
    voltage = np.asarray(run.pressure.valve_relay_v, dtype=float)
    time = np.asarray(run.pressure.time_s, dtype=float)
    finite = np.isfinite(voltage) & np.isfinite(time)
    if int(np.count_nonzero(finite)) < 2:
        raise ValueError("relay history has fewer than two finite samples")
    lower, upper = np.percentile(voltage[finite], [5.0, 95.0])
    if not upper > lower:
        raise ValueError("relay history has no resolvable open/closed amplitude")
    threshold = float(lower + fraction * (upper - lower))
    open_samples = finite & (voltage >= threshold)
    best_start = best_stop = start = 0
    current = 0
    for index, open_now in enumerate(open_samples):
        if open_now:
            if current == 0:
                start = index
            current += 1
            if current > best_stop - best_start:
                best_start, best_stop = start, index + 1
        else:
            current = 0
    if best_stop - best_start < 2:
        raise ValueError("relay history has no sustained open interval")
    return E31ValveInterval(
        start_s=float(time[best_start]),
        end_s=float(time[best_stop - 1]),
        threshold_v=threshold,
        samples=best_stop - best_start,
    )


def nozzle_pressure_rise(
    run: E31HighPressureRun | E31PressureRun,
    relay_interval: E31ValveInterval,
    *,
    required_rise_bar: float,
) -> E31NozzlePressureRise:
    """Locate the first PNoz response using an explicitly declared rise.

    ``required_rise_bar`` must be selected independently of the model result.
    The baseline is the median valid PNoz record before the relay-opening
    sample.  A direct pressure crossing avoids treating the electrical relay
    as the physical beginning of a jet, but it does not infer a valve area,
    mass flow, or a release end time.
    """
    rise = float(required_rise_bar)
    if not np.isfinite(rise) or rise <= 0.0:
        raise ValueError("required_rise_bar must be positive and finite")
    time = np.asarray(run.pressure.time_s, dtype=float)
    pressure = np.asarray(run.pressure.nozzle_pressure_bar, dtype=float)
    valid = np.isfinite(time) & np.isfinite(pressure) & (pressure > 0.0)
    baseline_samples = pressure[valid & (time < relay_interval.start_s)]
    if baseline_samples.size < 1:
        raise ValueError("PNoz has no pre-relay baseline sample")
    baseline = float(np.median(baseline_samples))
    threshold = baseline + rise
    candidates = np.flatnonzero(
        valid
        & (time >= relay_interval.start_s)
        & (time <= relay_interval.end_s)
        & (pressure >= threshold)
    )
    if candidates.size == 0:
        raise ValueError("PNoz does not cross the declared response threshold")
    response_time = float(time[int(candidates[0])])
    return E31NozzlePressureRise(
        relay_interval=relay_interval,
        baseline_pressure_bar=baseline,
        required_rise_bar=rise,
        threshold_pressure_bar=threshold,
        response_time_s=response_time,
        delay_after_relay_s=response_time - relay_interval.start_s,
    )


def axis_temperature_envelope(
    run: E31HighPressureRun,
    interval: E31ValveInterval,
) -> E31AxisTemperatureEnvelope:
    """Reduce raw T5/T6/T7 to a fixed interval without calibration correction.

    A reported minimum is retained to reveal the observed cold excursion; the
    5th percentile is the robust low-temperature comparison quantity; and
    the median records the warm occupancy.  The caller supplies the interval
    explicitly, so the reader never tunes a timing window against a model.
    """
    if not interval.end_s > interval.start_s:
        raise ValueError("temperature interval must have positive duration")
    history = run.axis_temperature
    mask = (
        np.isfinite(history.time_s)
        & (history.time_s >= interval.start_s)
        & (history.time_s <= interval.end_s)
    )
    if int(np.count_nonzero(mask)) < 2:
        raise ValueError("temperature history has fewer than two samples in interval")

    def summaries(values: np.ndarray) -> tuple[float, float, float]:
        chosen = np.asarray(values, dtype=float)[mask]
        chosen = chosen[np.isfinite(chosen)]
        if chosen.size < 2:
            raise ValueError("an axis thermocouple has fewer than two finite samples")
        return (
            float(np.min(chosen)),
            float(np.percentile(chosen, 5.0)),
            float(np.median(chosen)),
        )

    t5, t6, t7 = (summaries(values) for values in (history.t5_k, history.t6_k, history.t7_k))
    return E31AxisTemperatureEnvelope(
        interval=interval,
        samples=int(np.count_nonzero(mask)),
        minimum_k=(t5[0], t6[0], t7[0]),
        percentile_05_k=(t5[1], t6[1], t7[1]),
        median_k=(t5[2], t6[2], t7[2]),
    )


def ideal_choked_hydrogen_source_bound(
    run: E31HighPressureRun,
    *,
    storage_temperature_k: float,
    ambient_pressure_pa: float = 101325.0,
) -> E31IdealChokedSourceBound:
    """Map an E3.1 P/T/diameter record to an ideal pure-H2 source upper bound.

    ``storage_temperature_k`` is deliberately explicit because E3.1's closed
    thermocouples have a reported cold-bath bias. The result assumes a fully
    open sharp 4-mm-or-equivalent orifice (``Cd=1``). Any valve loss, opening
    history, non-pure composition or measured sub-unity discharge coefficient
    makes the true flow lower; this function does not choose one.
    """
    import math

    from ..addons.notional import (
        energy_conserving_notional_nozzle,
        isentropic_throat,
    )

    storage_temperature_k = float(storage_temperature_k)
    ambient_pressure_pa = float(ambient_pressure_pa)
    if storage_temperature_k <= 0.0 or ambient_pressure_pa <= 0.0:
        raise ValueError("source-bound temperatures and pressures must be positive")
    storage_pressure_pa = run.initial_vessel_pressure_bar * 1.0e5
    diameter_m = run.nozzle_diameter_mm * 1.0e-3
    throat = isentropic_throat(
        fluid="Hydrogen",
        storage_temperature=storage_temperature_k,
        storage_pressure=storage_pressure_pa,
        ambient_pressure=ambient_pressure_pa,
        allow_supercritical_gas=True,
    )
    mass_flow = throat.mass_flux * math.pi * diameter_m**2 / 4.0
    nozzle = energy_conserving_notional_nozzle(
        fluid="Hydrogen",
        storage_temperature=storage_temperature_k,
        storage_pressure=storage_pressure_pa,
        mass_flow=mass_flow,
        orifice_diameter=diameter_m,
        ambient_pressure=ambient_pressure_pa,
        allow_supercritical_gas=True,
    )
    if not nozzle.admissible:
        raise RuntimeError("ideal E3.1 source did not close at atmospheric pressure")
    assert nozzle.diameter is not None
    assert nozzle.temperature is not None
    assert nozzle.density is not None
    return E31IdealChokedSourceBound(
        storage_temperature_k=storage_temperature_k,
        storage_pressure_bar=run.initial_vessel_pressure_bar,
        orifice_diameter_mm=run.nozzle_diameter_mm,
        ideal_mass_flow_kg_s=mass_flow,
        throat_pressure_bar=throat.pressure / 1.0e5,
        throat_temperature_k=throat.temperature,
        throat_velocity_m_s=throat.velocity,
        atmospheric_diameter_m=nozzle.diameter,
        atmospheric_temperature_k=nozzle.temperature,
        atmospheric_density_kg_m3=nozzle.density,
        atmospheric_velocity_m_s=nozzle.velocity,
    )


def homogeneous_equilibrium_source_bound(
    run: E31HighPressureRun,
    *,
    storage_temperature_k: float,
    ambient_temperature_k: float = 293.0,
    ambient_pressure_pa: float = 101325.0,
) -> Any:
    """Return the fast phase-safe E3.1 source bound after H2 evaporation.

    This chains the explicit ``Cd=1`` pure-H2 source upper bound through its
    atmospheric liquid/vapour flash and the existing common-velocity,
    species-resolved N2/O2 evaporation endpoint.  It is an equilibrium
    *lower-distance* bound: it cannot establish a finite droplet lifetime or
    validate gas/particle slip.  The returned source has sufficient ambient
    dilution that its H2 partial pressure is below saturation, so it is a
    physically admissible handoff candidate for a gas-jet calculation.
    """
    from ..addons.lh2_droplets import (
        flashing_hydrogen_droplet_source,
        homogeneous_equilibrium_hydrogen_source_from_postflash,
    )

    bound = ideal_choked_hydrogen_source_bound(
        run,
        storage_temperature_k=storage_temperature_k,
        ambient_pressure_pa=ambient_pressure_pa,
    )
    postflash = flashing_hydrogen_droplet_source(
        mass_flow=bound.ideal_mass_flow_kg_s,
        orifice_diameter=run.nozzle_diameter_mm * 1.0e-3,
        upstream_temperature=storage_temperature_k,
        upstream_pressure=run.initial_vessel_pressure_bar * 1.0e5,
        ambient_temperature=ambient_temperature_k,
        ambient_pressure=ambient_pressure_pa,
    )
    return homogeneous_equilibrium_hydrogen_source_from_postflash(
        postflash,
        ambient_temperature=ambient_temperature_k,
        ambient_pressure=ambient_pressure_pa,
    )


def compare_axis_temperature_envelope(
    envelope: E31AxisTemperatureEnvelope,
    model_temperature_k: tuple[float, float, float] | list[float] | np.ndarray,
) -> E31AxisTemperatureComparison:
    """Report a fixed three-station comparison without selecting a time peak.

    ``model_temperature_k`` must correspond, in order, to T5/T6/T7 at the
    documented 0.25/0.75/1.75-m axial locations.  The residual is observation
    5th-percentile minus model temperature: a positive result means that the
    stationary model is colder than that robust observed low excursion.
    """
    values = tuple(float(value) for value in model_temperature_k)
    if len(values) != 3 or not all(np.isfinite(value) and value > 0.0 for value in values):
        raise ValueError("supply three positive finite T5/T6/T7 model temperatures")
    residual = tuple(
        observed - model
        for observed, model in zip(envelope.percentile_05_k, values)
    )
    inside = tuple(
        low <= model <= high
        for low, high, model in zip(
            envelope.minimum_k, envelope.median_k, values
        )
    )
    return E31AxisTemperatureComparison(
        distances_m=(0.25, 0.75, 1.75),
        model_temperature_k=values,
        observed_minimum_k=envelope.minimum_k,
        observed_percentile_05_k=envelope.percentile_05_k,
        observed_median_k=envelope.median_k,
        residual_to_percentile_05_k=residual,
        inside_recorded_minimum_median=inside,
    )


def compare_blowdown_pressure_stations(
    run: E31HighPressureRun | E31PressureRun,
    *,
    source_start_time_s: float,
    model_time_s: np.ndarray | list[float] | tuple[float, ...],
    model_pressure_pa: np.ndarray | list[float] | tuple[float, ...],
    model_tank_temperature_k: np.ndarray | list[float] | tuple[float, ...],
    stations_after_source_start_s: np.ndarray | list[float] | tuple[float, ...],
) -> E31BlowdownPressureComparison:
    """Compare explicit pressure stations without fitting an opening delay.

    ``source_start_time_s`` is normally a declared PNoz response time, not
    the electrical relay time.  The function linearly interpolates only the
    measured pressure trace; requested stations outside either record are
    rejected rather than extrapolated or moved to a better-looking time.
    """
    start = float(source_start_time_s)
    model_time = np.asarray(model_time_s, dtype=float)
    model_pressure = np.asarray(model_pressure_pa, dtype=float)
    model_temperature = np.asarray(model_tank_temperature_k, dtype=float)
    stations = np.asarray(stations_after_source_start_s, dtype=float)
    if not np.isfinite(start):
        raise ValueError("source_start_time_s must be finite")
    if model_time.ndim != 1 or model_time.size < 2:
        raise ValueError("model_time_s must contain at least two values")
    if not (model_pressure.shape == model_time.shape == model_temperature.shape):
        raise ValueError("model time, pressure and temperature shapes must match")
    if not np.all(np.isfinite(model_time)) or not np.all(np.diff(model_time) > 0.0):
        raise ValueError("model_time_s must be finite and strictly increasing")
    if not np.all(np.isfinite(model_pressure)) or np.any(model_pressure <= 0.0):
        raise ValueError("model_pressure_pa must be positive and finite")
    if not np.all(np.isfinite(model_temperature)) or np.any(model_temperature <= 0.0):
        raise ValueError("model_tank_temperature_k must be positive and finite")
    if stations.ndim != 1 or stations.size == 0 or not np.all(np.isfinite(stations)):
        raise ValueError("stations_after_source_start_s must be a finite non-empty vector")
    measured_time = np.asarray(run.pressure.time_s, dtype=float)
    measured_pressure = np.asarray(run.pressure.vessel_pressure_bar, dtype=float)
    valid = np.isfinite(measured_time) & np.isfinite(measured_pressure) & (measured_pressure > 0.0)
    measured_time, measured_pressure = measured_time[valid], measured_pressure[valid]
    order = np.argsort(measured_time)
    measured_time, measured_pressure = measured_time[order], measured_pressure[order]
    requested_time = start + stations
    if (
        requested_time.min() < measured_time[0]
        or requested_time.max() > measured_time[-1]
        or stations.min() < model_time[0]
        or stations.max() > model_time[-1]
    ):
        raise ValueError("requested pressure station lies outside a supplied record")
    observed = np.interp(requested_time, measured_time, measured_pressure)
    predicted = np.interp(stations, model_time, model_pressure) / 1.0e5
    predicted_temperature = np.interp(stations, model_time, model_temperature)
    return E31BlowdownPressureComparison(
        source_start_time_s=start,
        time_after_source_start_s=tuple(float(value) for value in stations),
        measured_pressure_bar=tuple(float(value) for value in observed),
        model_pressure_bar=tuple(float(value) for value in predicted),
        pressure_residual_bar=tuple(float(value) for value in predicted - observed),
        model_tank_temperature_k=tuple(float(value) for value in predicted_temperature),
    )


def vessel_inventory_change(
    run: E31HighPressureRun,
    interval: E31ValveInterval,
    *,
    vessel_volume_m3: float,
    storage_temperature_k: float,
) -> E31VesselInventoryChange:
    """Infer integrated released H2 mass from a declared vessel volume and T.

    Pressure in the E3.1 ``Press`` sheet is used as an absolute pure-H2 state
    variable. ``storage_temperature_k`` is explicit because the campaign's
    closed thermocouples show a documented cold-bath bias; callers should
    report a temperature sensitivity rather than silently choosing a
    correction.  The function returns only interval-integrated mass loss, not
    a noisy differentiated mass-flow trace.
    """
    from CoolProp.CoolProp import PropsSI

    volume = float(vessel_volume_m3)
    temperature = float(storage_temperature_k)
    if not np.isfinite(volume) or volume <= 0.0:
        raise ValueError("vessel_volume_m3 must be positive and finite")
    if not np.isfinite(temperature) or temperature <= 0.0:
        raise ValueError("storage_temperature_k must be positive and finite")
    if not interval.end_s > interval.start_s:
        raise ValueError("inventory interval must have positive duration")
    _time, mass = _vessel_mass_series(
        run, interval, vessel_volume_m3=volume, storage_temperature_k=temperature
    )
    initial_mass = float(mass[0])
    final_mass = float(mass[-1])
    released = initial_mass - final_mass
    if released < -1.0e-10:
        raise ValueError("vessel mass rises over the declared release interval")
    released = max(released, 0.0)
    return E31VesselInventoryChange(
        interval=interval,
        vessel_volume_m3=volume,
        storage_temperature_k=temperature,
        samples=int(mass.size),
        initial_mass_kg=initial_mass,
        final_mass_kg=final_mass,
        released_mass_kg=released,
        mean_release_rate_kg_s=released / (interval.end_s - interval.start_s),
    )


def _vessel_mass_series(
    run: E31HighPressureRun,
    interval: E31ValveInterval,
    *,
    vessel_volume_m3: float,
    storage_temperature_k: float,
) -> tuple[np.ndarray, np.ndarray]:
    """Return interval time and pure-H2 inventory after common validation."""
    from CoolProp.CoolProp import PropsSI

    time = np.asarray(run.pressure.time_s, dtype=float)
    pressure_bar = np.asarray(run.pressure.vessel_pressure_bar, dtype=float)
    mask = (
        np.isfinite(time)
        & np.isfinite(pressure_bar)
        & (pressure_bar > 0.0)
        & (time >= interval.start_s)
        & (time <= interval.end_s)
    )
    if int(np.count_nonzero(mask)) < 2:
        raise ValueError("vessel pressure has fewer than two valid interval samples")
    density = np.asarray(
        PropsSI(
            "D", "T", storage_temperature_k,
            "P", pressure_bar[mask] * 1.0e5, "Hydrogen",
        ),
        dtype=float,
    )
    if not np.all(np.isfinite(density)) or np.any(density <= 0.0):
        raise ValueError("vessel EOS returned a non-positive density")
    return time[mask], density * vessel_volume_m3


def vessel_inventory_release_timing(
    run: E31HighPressureRun,
    interval: E31ValveInterval,
    *,
    vessel_volume_m3: float,
    storage_temperature_k: float,
    released_fraction: float = 0.95,
) -> E31InventoryReleaseTiming:
    """Locate a cumulative inventory-loss fraction without differentiating P.

    The cumulative loss is made non-decreasing only to suppress late
    instrument noise after blowdown; this does not alter the early pressure
    record.  It exposes discharge timing but does not construct an
    instantaneous mass-flow trace.
    """
    fraction = float(released_fraction)
    if not 0.0 < fraction <= 1.0:
        raise ValueError("released_fraction must lie in (0, 1]")
    volume = float(vessel_volume_m3)
    temperature = float(storage_temperature_k)
    if not np.isfinite(volume) or volume <= 0.0:
        raise ValueError("vessel_volume_m3 must be positive and finite")
    if not np.isfinite(temperature) or temperature <= 0.0:
        raise ValueError("storage_temperature_k must be positive and finite")
    time, mass = _vessel_mass_series(
        run, interval, vessel_volume_m3=volume, storage_temperature_k=temperature
    )
    cumulative = np.maximum.accumulate(np.maximum(mass[0] - mass, 0.0))
    total = float(cumulative[-1])
    if total <= 0.0:
        raise ValueError("vessel inventory does not decrease over the interval")
    index = int(np.searchsorted(cumulative, fraction * total, side="left"))
    index = min(index, len(time) - 1)
    return E31InventoryReleaseTiming(
        interval=interval,
        storage_temperature_k=temperature,
        vessel_volume_m3=volume,
        released_fraction=fraction,
        time_after_open_s=float(time[index] - interval.start_s),
        released_mass_kg=float(cumulative[index]),
        total_released_mass_kg=total,
    )
