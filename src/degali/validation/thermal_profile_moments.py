"""Coefficient-free moments of paired PRESLHY temperature/H2 profiles.

These helpers define an observation operator.  They neither select a scalar
diffusivity nor alter a dispersion-model state.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path

import numpy as np

from .preslhy import sensor_positions


_TEMPERATURE = re.compile(
    r"^(Centreline|Above1|Above2|Below1|Below2)_"
    r"(\d{3})_TC\d+C$"
)
_XSENSOR = re.compile(r"_(\d{2}[A-Z]{2}\d+)Output$")
_VERTICAL_OFFSETS = {
    "Below2": -0.50,
    "Below1": -0.25,
    "Centreline": 0.0,
    "Above1": 0.25,
    "Above2": 0.50,
}
_STATIONS = (1.78, 4.00)
_OFFSETS = np.asarray((-0.50, -0.25, 0.0, 0.25, 0.50))


@dataclass(frozen=True)
class TruncatedProfileMoment:
    zeroth: float
    centre: float
    variance: float

    @property
    def width(self) -> float:
        return math.sqrt(self.variance)


@dataclass(frozen=True)
class LagEstimate:
    seconds: int
    correlation: float
    samples: int


@dataclass(frozen=True)
class ModelProfileProjection:
    station: float
    release_height: float
    coordinate: np.ndarray
    absolute_height: np.ndarray
    ambient_temperature: float
    temperature: np.ndarray
    thermal_deficit: np.ndarray
    hydrogen: np.ndarray
    thermal_moment: TruncatedProfileMoment
    hydrogen_moment: TruncatedProfileMoment


def truncated_profile_moment(coordinate, scalar) -> TruncatedProfileMoment:
    """Return zeroth, centroid and central second moment on a finite line."""
    z = np.asarray(coordinate, dtype=float)
    value = np.asarray(scalar, dtype=float)
    if z.ndim != 1 or value.shape != z.shape or z.size < 3:
        raise ValueError("matching one-dimensional profile with at least 3 points required")
    if not np.all(np.isfinite(z)) or not np.all(np.isfinite(value)):
        raise ValueError("finite profile coordinates and values required")
    if np.any(np.diff(z) <= 0.0) or np.any(value < 0.0):
        raise ValueError("strictly increasing coordinates and nonnegative scalar required")
    zeroth = float(np.trapezoid(value, z))
    if zeroth <= 0.0:
        raise ValueError("positive truncated zeroth moment required")
    centre = float(np.trapezoid(z * value, z) / zeroth)
    variance = float(np.trapezoid((z - centre) ** 2 * value, z) / zeroth)
    if variance < -1.0e-14:
        raise ValueError("negative profile variance")
    return TruncatedProfileMoment(zeroth, centre, max(variance, 0.0))


def project_model_profile(
    trajectory,
    station: float,
    release_height: float,
    *,
    coordinate=_OFFSETS,
    thermal_centre_gate: float = 5.0,
    hydrogen_centre_gate: float = 1.0,
) -> ModelProfileProjection:
    """Sample a steady trajectory with the finite observed-profile operator.

    The supplied coordinates are relative to the experimental release axis,
    while the trajectory receptor API uses absolute height above ground.  This
    conversion is intentionally part of the observation operator.
    """
    station = float(station)
    release_height = float(release_height)
    z = np.asarray(coordinate, dtype=float)
    gates = (float(thermal_centre_gate), float(hydrogen_centre_gate))
    if (
        not math.isfinite(station)
        or station <= 0.0
        or not math.isfinite(release_height)
        or release_height < 0.0
        or z.ndim != 1
        or z.size < 3
        or not np.all(np.isfinite(z))
        or np.any(np.diff(z) <= 0.0)
        or not all(math.isfinite(gate) and gate >= 0.0 for gate in gates)
    ):
        raise ValueError("finite station, release height, coordinates and gates required")
    centre = int(np.argmin(np.abs(z)))
    if abs(z[centre]) > 1.0e-12:
        raise ValueError("profile coordinates must contain the release-axis centre")
    if not hasattr(trajectory, "state_at") or trajectory.state_at(station) is None:
        raise ValueError("trajectory does not cover the requested station")
    try:
        ambient = float(trajectory.model.thermodynamics.ambient_temperature)
    except (AttributeError, TypeError, ValueError) as error:
        raise ValueError("trajectory must expose a finite ambient temperature") from error
    absolute = release_height + z
    if np.any(absolute < 0.0):
        raise ValueError("model receptor heights cannot be below ground")
    temperature = np.asarray([
        trajectory.temperature_at(station, 0.0, height) for height in absolute
    ], dtype=float)
    hydrogen = np.asarray([
        trajectory.concentration_at(station, 0.0, height) for height in absolute
    ], dtype=float)
    if (
        not math.isfinite(ambient)
        or ambient <= 0.0
        or not np.all(np.isfinite(temperature))
        or np.any(temperature <= 0.0)
        or not np.all(np.isfinite(hydrogen))
        or np.any(hydrogen < 0.0)
        or np.any(hydrogen > 100.0)
    ):
        raise ValueError("model receptor profile is outside its physical range")
    thermal = np.maximum(ambient - temperature, 0.0)
    if thermal[centre] < gates[0] or hydrogen[centre] < gates[1]:
        raise ValueError("model profile does not pass the frozen centre signal gates")
    return ModelProfileProjection(
        station=station,
        release_height=release_height,
        coordinate=z.copy(),
        absolute_height=absolute,
        ambient_temperature=ambient,
        temperature=temperature,
        thermal_deficit=thermal,
        hydrogen=hydrogen,
        thermal_moment=truncated_profile_moment(z, thermal),
        hydrogen_moment=truncated_profile_moment(z, hydrogen),
    )


def best_integer_lag(upstream, downstream, *, maximum_seconds: int = 5) -> LagEstimate:
    """Find the nonnegative integer lag maximizing ordinary correlation."""
    up = np.asarray(upstream, dtype=float)
    down = np.asarray(downstream, dtype=float)
    if up.ndim != 1 or down.shape != up.shape or up.size < maximum_seconds + 3:
        raise ValueError("matching series longer than the lag window required")
    if not np.all(np.isfinite(up)) or not np.all(np.isfinite(down)):
        raise ValueError("finite lag series required")
    if not isinstance(maximum_seconds, int) or maximum_seconds < 0:
        raise ValueError("nonnegative integer maximum lag required")
    candidates = []
    for lag in range(maximum_seconds + 1):
        left = up if lag == 0 else up[:-lag]
        right = down if lag == 0 else down[lag:]
        if np.std(left) == 0.0 or np.std(right) == 0.0:
            correlation = -math.inf
        else:
            correlation = float(np.corrcoef(left, right)[0, 1])
        candidates.append((correlation, -lag, len(left)))
    correlation, negative_lag, samples = max(candidates)
    if not math.isfinite(correlation):
        raise ValueError("lag correlation is undefined for constant series")
    return LagEstimate(-negative_lag, correlation, samples)


def _clock_seconds(value) -> float:
    if not isinstance(value, str):
        raise ValueError("PRESLHY clock must be a string")
    fields = value.strip().split(":")
    if len(fields) != 3:
        raise ValueError(f"unrecognised PRESLHY clock {value!r}")
    hour, minute, second = map(float, fields)
    return 3600.0 * hour + 60.0 * minute + second


def _column(rows, index: int) -> np.ndarray:
    values = []
    for row in rows:
        value = row[index] if index < len(row) else None
        try:
            number = float(value)
        except (TypeError, ValueError):
            number = math.nan
        values.append(number)
    return np.asarray(values)


def _interpolate_bracketed(source_time, source_value, target_time) -> np.ndarray:
    valid = np.isfinite(source_time) & np.isfinite(source_value)
    time = np.asarray(source_time)[valid]
    value = np.asarray(source_value)[valid]
    if time.size < 2 or np.any(np.diff(time) <= 0.0):
        raise ValueError("increasing source clock with at least two finite values required")
    target = np.asarray(target_time, dtype=float)
    out = np.interp(target, time, value)
    out[(target < time[0]) | (target > time[-1])] = np.nan
    return out


def read_paired_profiles(
    workbook: str | Path,
    report: str | Path,
    *,
    flexlogger_rows: tuple[int, int],
) -> dict:
    """Read simultaneous vertical temperature-deficit and H2 profiles.

    ``flexlogger_rows`` is one-based and inclusive, matching Excel notation.
    The returned coordinate is relative to the release axis.
    """
    import openpyxl

    path = Path(workbook)
    first, last = flexlogger_rows
    if first < 2 or last < first:
        raise ValueError("valid inclusive Flexlogger data rows required")
    positions = sensor_positions(report)
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    flex = book["Flexlogger"]
    flex_header = list(next(flex.iter_rows(min_row=1, max_row=1, values_only=True)))
    flex_rows = list(flex.iter_rows(min_row=first, max_row=last, values_only=True))
    target_time = np.asarray([_clock_seconds(row[0]) for row in flex_rows])
    if not np.allclose(np.diff(target_time), 1.0, rtol=0.0, atol=1.0e-8):
        raise ValueError("selected Flexlogger window is not uniformly sampled at 1 Hz")
    ambient_index = flex_header.index("Ambient_Temperature")
    ambient = _column(flex_rows, ambient_index) + 273.15

    temperature_columns = {}
    for index, heading in enumerate(flex_header):
        match = _TEMPERATURE.match(str(heading))
        if match is None:
            continue
        placement, distance = match.groups()
        station = float(distance) / 100.0
        if station in _STATIONS:
            temperature_columns[(station, _VERTICAL_OFFSETS[placement])] = index

    xensor = book["Xensor"]
    x_rows = list(xensor.iter_rows(values_only=True))
    x_header, x_data = list(x_rows[0]), x_rows[1:]
    source_time = np.asarray([_clock_seconds(row[0]) for row in x_data])
    concentration_columns = {}
    for index, heading in enumerate(x_header):
        match = _XSENSOR.search(str(heading))
        if match is None or match.group(1) not in positions:
            continue
        x, y, z_axis = positions[match.group(1)]
        station = next(
            (candidate for candidate in _STATIONS if abs(x - candidate) < 1.0e-12),
            None,
        )
        offset = z_axis - 0.5
        matched_offset = next(
            (candidate for candidate in _OFFSETS if abs(offset - candidate) < 1.0e-12),
            None,
        )
        if station is not None and abs(y) < 1.0e-12 and matched_offset is not None:
            concentration_columns[(station, matched_offset)] = index

    expected = {(station, offset) for station in _STATIONS for offset in _OFFSETS}
    if set(temperature_columns) != expected or set(concentration_columns) != expected:
        missing_t = sorted(expected - set(temperature_columns))
        missing_c = sorted(expected - set(concentration_columns))
        raise ValueError(f"incomplete collocated cross: temperature={missing_t}, H2={missing_c}")

    profiles = {}
    clipped = dict(thermal=0, hydrogen_low=0, hydrogen_high=0)
    for station in _STATIONS:
        temperature = np.column_stack([
            _column(flex_rows, temperature_columns[(station, offset)]) + 273.15
            for offset in _OFFSETS
        ])
        concentration = np.column_stack([
            _interpolate_bracketed(
                source_time,
                _column(x_data, concentration_columns[(station, offset)]),
                target_time,
            )
            for offset in _OFFSETS
        ])
        if not np.all(np.isfinite(temperature)) or not np.all(np.isfinite(concentration)):
            raise ValueError("selected paired profiles contain unbracketed or nonfinite values")
        deficit_raw = ambient[:, None] - temperature
        clipped["thermal"] += int(np.sum(deficit_raw < 0.0))
        clipped["hydrogen_low"] += int(np.sum(concentration < 0.0))
        clipped["hydrogen_high"] += int(np.sum(concentration > 100.0))
        profiles[str(station)] = dict(
            thermal_deficit=np.maximum(deficit_raw, 0.0),
            hydrogen=np.clip(concentration, 0.0, 100.0),
            temperature=temperature,
        )
    book.close()
    return dict(
        time_seconds=target_time,
        ambient_temperature=ambient,
        coordinate=_OFFSETS.copy(),
        profiles=profiles,
        clipped_values=clipped,
        flexlogger_rows=[first, last],
    )


def summarize_paired_profile_moments(data: dict) -> dict:
    """Apply the pre-registered signal, lag and truncated-moment rules."""
    z = np.asarray(data["coordinate"], dtype=float)
    profile_moments = {}
    for station in _STATIONS:
        entry = data["profiles"][str(station)]
        station_result = {}
        for name, threshold in (("thermal_deficit", 5.0), ("hydrogen", 1.0)):
            values = np.asarray(entry[name], dtype=float)
            rows = []
            for index, value in enumerate(values):
                usable = bool(value[2] >= threshold)
                moment = truncated_profile_moment(z, value) if usable else None
                rows.append(dict(
                    usable=usable,
                    zeroth=None if moment is None else moment.zeroth,
                    centre=None if moment is None else moment.centre,
                    variance=None if moment is None else moment.variance,
                    width=None if moment is None else moment.width,
                ))
            station_result[name] = rows
        profile_moments[str(station)] = station_result

    lag = {}
    for name in ("thermal_deficit", "hydrogen"):
        upstream = data["profiles"][str(_STATIONS[0])][name][:, 2]
        downstream = data["profiles"][str(_STATIONS[1])][name][:, 2]
        estimate = best_integer_lag(upstream, downstream, maximum_seconds=5)
        lag[name] = dict(
            seconds=estimate.seconds,
            correlation=estimate.correlation,
            samples=estimate.samples,
        )
    common_delay_passed = bool(
        abs(lag["thermal_deficit"]["seconds"] - lag["hydrogen"]["seconds"]) <= 1
        and min(lag["thermal_deficit"]["correlation"], lag["hydrogen"]["correlation"]) >= 0.5
    )
    common_lag = int(math.floor(
        0.5 * (lag["thermal_deficit"]["seconds"] + lag["hydrogen"]["seconds"]) + 0.5
    ))
    pairs = []
    count = len(data["time_seconds"]) - common_lag
    for index in range(max(count, 0)):
        downstream_index = index + common_lag
        row = dict(
            upstream_time_seconds=float(data["time_seconds"][index]),
            downstream_time_seconds=float(data["time_seconds"][downstream_index]),
        )
        usable = True
        for name in ("thermal_deficit", "hydrogen"):
            up = profile_moments[str(_STATIONS[0])][name][index]
            down = profile_moments[str(_STATIONS[1])][name][downstream_index]
            usable = usable and up["usable"] and down["usable"]
            if up["usable"] and down["usable"]:
                row[name] = dict(
                    upstream_width=up["width"],
                    downstream_width=down["width"],
                    squared_width_growth=down["variance"] - up["variance"],
                )
        if usable:
            denominator = row["hydrogen"]["squared_width_growth"]
            row["growth_ratio"] = (
                None if abs(denominator) <= 1.0e-12
                else row["thermal_deficit"]["squared_width_growth"] / denominator
            )
            pairs.append(row)

    def finite(values):
        return np.asarray([value for value in values if value is not None and math.isfinite(value)])

    station_summary = {}
    for station in _STATIONS:
        out = {}
        for name in ("thermal_deficit", "hydrogen"):
            widths = finite([row["width"] for row in profile_moments[str(station)][name]])
            out[name] = dict(
                usable_profiles=int(widths.size),
                median_width=float(np.median(widths)) if widths.size else None,
                p25_width=float(np.percentile(widths, 25)) if widths.size else None,
                p75_width=float(np.percentile(widths, 75)) if widths.size else None,
            )
        ratios = []
        thermal = profile_moments[str(station)]["thermal_deficit"]
        hydrogen = profile_moments[str(station)]["hydrogen"]
        for t, c in zip(thermal, hydrogen):
            if t["usable"] and c["usable"] and c["width"] > 0.0:
                ratios.append(t["width"] / c["width"])
        ratio = finite(ratios)
        out["simultaneous_thermal_to_species_width_ratio"] = dict(
            count=int(ratio.size),
            median=float(np.median(ratio)) if ratio.size else None,
            p25=float(np.percentile(ratio, 25)) if ratio.size else None,
            p75=float(np.percentile(ratio, 75)) if ratio.size else None,
        )
        station_summary[str(station)] = out

    ratios = finite([row["growth_ratio"] for row in pairs])
    thermal_growth = finite([row["thermal_deficit"]["squared_width_growth"] for row in pairs])
    species_growth = finite([row["hydrogen"]["squared_width_growth"] for row in pairs])
    enough_pairs = len(pairs) >= 10
    inference_issued = common_delay_passed and enough_pairs
    return dict(
        lag=lag,
        common_lag_seconds=common_lag,
        common_delay_passed=common_delay_passed,
        enough_pairs=enough_pairs,
        inference_issued=inference_issued,
        paired_profiles=len(pairs),
        station_summary=station_summary,
        growth_summary=dict(
            ratios=int(ratios.size),
            median_ratio=float(np.median(ratios)) if ratios.size else None,
            p25_ratio=float(np.percentile(ratios, 25)) if ratios.size else None,
            p75_ratio=float(np.percentile(ratios, 75)) if ratios.size else None,
            thermal_positive=int(np.sum(thermal_growth > 0.0)),
            thermal_nonpositive=int(np.sum(thermal_growth <= 0.0)),
            species_positive=int(np.sum(species_growth > 0.0)),
            species_nonpositive=int(np.sum(species_growth <= 0.0)),
        ),
        pairs=pairs,
        profile_moments=profile_moments,
    )


__all__ = [
    "LagEstimate",
    "ModelProfileProjection",
    "TruncatedProfileMoment",
    "best_integer_lag",
    "project_model_profile",
    "read_paired_profiles",
    "summarize_paired_profile_moments",
    "truncated_profile_moment",
]
