"""Read the public HyTunnel-CS Pro-Science concentration/flow matrices.

The record at DOI ``10.5281/zenodo.4668554`` provides four concentration
workbooks, one ventilation-flow workbook and the Inventor CAD assembly.  The
workbooks use repeated x/y matrix blocks for average, maximum and minimum
values.  This intake preserves those published statistics without inferring
an unavailable transient clock or measurement height.
"""

from __future__ import annotations

from dataclasses import dataclass
import math
from pathlib import Path
import re
from typing import Literal


DATASET_DOI = "10.5281/zenodo.4668554"
Statistic = Literal["average", "maximum", "minimum"]


def _openpyxl():
    try:
        import openpyxl
    except ImportError as error:  # pragma: no cover - environment boundary
        raise RuntimeError(
            "HyTunnel workbook intake requires openpyxl; install DEGALI's "
            "test/research dependencies"
        ) from error
    return openpyxl


def _normal(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower()).strip("_")


def _finite_number(value: object) -> float | None:
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _condition_from_title(title: str) -> tuple[float | None, float | None, float, str, str]:
    nozzle_match = re.search(r"(?P<value>\d+(?:\.\d+)?)\s*mm", title, re.IGNORECASE)
    mass_match = re.search(r"(?P<value>\d+(?:\.\d+)?)\s*g\b", title, re.IGNORECASE)
    wind_match = re.search(r"W\s*=\s*(?P<value>\d+(?:\.\d+)?)", title, re.IGNORECASE)
    if wind_match is None and any(token in title.lower() for token in ("co-flow", "counter", "cross")):
        # One published sheet omits the ``W`` in ``H2 = 5 counter-flow``.
        # Limit this fallback to directional titles so that gas/mass-flow
        # numbers are not mistaken for a ventilation speed.
        wind_match = re.search(
            r"=\s*(?P<value>1\.5|3\.5|5(?:\.0)?)\s*(?:m\s*s)?\s*"
            r"(?:co-flow|counter|cross)",
            title,
            re.IGNORECASE,
        )
    if wind_match is None:
        wind_match = re.search(
            r"(?P<value>1\.5|3\.5|5(?:\.0)?)\)?\s*m\s*(?:pro|/)?\s*s",
            title, re.IGNORECASE,
        )
    lower = title.lower()
    direction = "quiescent"
    if "counter" in lower:
        direction = "counter-flow"
    elif "cross" in lower:
        direction = "cross-flow"
    elif "co-flow" in lower:
        direction = "co-flow"
    gas = "He" if re.search(r"\bHe\b", title) else "H2"
    return (
        None if nozzle_match is None else float(nozzle_match.group("value")),
        None if mass_match is None else float(mass_match.group("value")),
        0.0 if wind_match is None else float(wind_match.group("value")),
        direction,
        gas,
    )


def _statistic_label(value: object) -> Statistic | None:
    text = str(value).lower() if value is not None else ""
    if "average" in text:
        return "average"
    if "maximum" in text or re.search(r"\bmax\b", text):
        return "maximum"
    if "minimum" in text or re.search(r"\bmin\b", text):
        return "minimum"
    return None


def _matrix_headers(rows: list[tuple[object, ...]]) -> list[tuple[int, int, Statistic]]:
    locations: list[tuple[int, int]] = []
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            if _normal(value) in {"x_mm", "x_m_m"}:
                locations.append((row_index, column_index))
    columns = sorted({column for _, column in locations})
    default = {
        column: statistic
        for column, statistic in zip(columns, ("average", "maximum", "minimum"))
    }
    headers: list[tuple[int, int, Statistic]] = []
    for row_index, column_index in locations:
        statistic = None
        for prior in range(row_index - 1, -1, -1):
            for nearby in range(max(0, column_index - 1), column_index + 3):
                if nearby >= len(rows[prior]):
                    continue
                statistic = _statistic_label(rows[prior][nearby])
                if statistic is not None:
                    break
            if statistic is not None:
                break
        statistic = statistic or default.get(column_index)
        if statistic is None:
            raise ValueError("HyTunnel matrix statistic could not be identified")
        headers.append((row_index, column_index, statistic))
    return headers


def _y_coordinates(
    rows: list[tuple[object, ...]], row_index: int, column_index: int,
) -> list[tuple[int, float]]:
    for candidate_row in (row_index, row_index - 1):
        if candidate_row < 0:
            continue
        result: list[tuple[int, float]] = []
        for column in range(column_index + 1, len(rows[candidate_row])):
            value = _finite_number(rows[candidate_row][column])
            if value is None:
                if result:
                    break
                continue
            result.append((column, value / 1000.0))
        if result:
            return result
    raise ValueError("HyTunnel matrix has no numerical y-coordinate headings")


def _grid_points(rows: list[tuple[object, ...]]):
    for header_row, x_column, statistic in _matrix_headers(rows):
        y_columns = _y_coordinates(rows, header_row, x_column)
        for row_index in range(header_row + 1, len(rows)):
            row = rows[row_index]
            if x_column >= len(row):
                break
            x_mm = _finite_number(row[x_column])
            if x_mm is None:
                break
            for y_column, y_m in y_columns:
                if y_column >= len(row):
                    continue
                value = _finite_number(row[y_column])
                if value is not None:
                    yield statistic, x_mm / 1000.0, y_m, value


@dataclass(frozen=True)
class HyTunnelConfiguration:
    configuration_id: str
    workbook: str
    sheet: str
    nozzle_diameter_mm: float | None
    nominal_mass_flow_g_s: float | None
    nominal_ventilation_speed_m_s: float
    flow_direction: str
    gas_label: str

    @property
    def is_direct_hydrogen(self) -> bool:
        return self.gas_label == "H2"


@dataclass(frozen=True)
class HyTunnelConcentrationPoint:
    configuration_id: str
    statistic: Statistic
    x_m: float
    y_m: float
    concentration_vol_percent: float


@dataclass(frozen=True)
class HyTunnelFlowPoint:
    configuration_id: str
    statistic: Statistic
    x_m: float
    y_m: float
    velocity_x_m_s: float


@dataclass(frozen=True)
class HyTunnelConcentrationDataset:
    configurations: tuple[HyTunnelConfiguration, ...]
    points: tuple[HyTunnelConcentrationPoint, ...]
    dataset_doi: str = DATASET_DOI

    @property
    def quantitative_lh2_validation_allowed(self) -> bool:
        return False

    @property
    def validation_scope(self) -> str:
        return "gh2_ventilated_jet_spatial_component_only"


@dataclass(frozen=True)
class HyTunnelFlowDataset:
    configuration_ids: tuple[str, ...]
    points: tuple[HyTunnelFlowPoint, ...]
    dataset_doi: str = DATASET_DOI


def read_hytunnel_concentration_workbook(
    path: str | Path,
) -> HyTunnelConcentrationDataset:
    """Read all average/maximum/minimum concentration matrices in one file."""

    openpyxl = _openpyxl()
    source = Path(path)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    configurations: list[HyTunnelConfiguration] = []
    points: list[HyTunnelConcentrationPoint] = []
    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
        configuration_id = f"{source.stem}:sheet-{sheet_index:02d}"
        nozzle, mass_flow, wind, direction, gas = _condition_from_title(worksheet.title)
        configurations.append(HyTunnelConfiguration(
            configuration_id=configuration_id,
            workbook=source.name,
            sheet=worksheet.title,
            nozzle_diameter_mm=nozzle,
            nominal_mass_flow_g_s=mass_flow,
            nominal_ventilation_speed_m_s=wind,
            flow_direction=direction,
            gas_label=gas,
        ))
        for statistic, x_m, y_m, concentration in _grid_points(rows):
            points.append(HyTunnelConcentrationPoint(
                configuration_id=configuration_id,
                statistic=statistic,
                x_m=x_m,
                y_m=y_m,
                concentration_vol_percent=concentration,
            ))
    if not configurations or not points:
        raise ValueError(f"HyTunnel workbook {source.name} contains no matrices")
    return HyTunnelConcentrationDataset(tuple(configurations), tuple(points))


def read_hytunnel_flow_workbook(path: str | Path) -> HyTunnelFlowDataset:
    """Read the one-minute x-velocity average/maximum/minimum matrices."""

    openpyxl = _openpyxl()
    source = Path(path)
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    ids: list[str] = []
    points: list[HyTunnelFlowPoint] = []
    for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
        rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
        configuration_id = f"{source.stem}:sheet-{sheet_index:02d}"
        ids.append(configuration_id)
        for statistic, x_m, y_m, velocity in _grid_points(rows):
            points.append(HyTunnelFlowPoint(
                configuration_id=configuration_id,
                statistic=statistic,
                x_m=x_m,
                y_m=y_m,
                velocity_x_m_s=velocity,
            ))
    if not ids or not points:
        raise ValueError(f"HyTunnel flow workbook {source.name} contains no matrices")
    return HyTunnelFlowDataset(tuple(ids), tuple(points))


__all__ = [
    "DATASET_DOI", "HyTunnelConcentrationDataset", "HyTunnelConcentrationPoint",
    "HyTunnelConfiguration", "HyTunnelFlowDataset", "HyTunnelFlowPoint",
    "read_hytunnel_concentration_workbook", "read_hytunnel_flow_workbook",
]
